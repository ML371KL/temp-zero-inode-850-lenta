# -*- coding: utf-8 -*-
"""Сборщики Ленты на фикстурах (D15): рынок, раскрытие, датабук, «Работа России».

В сеть тесты не ходят: `sources.fetch` подменён ответами из `tests/fixtures/`
(вежливые записи 28.09.2026, `tests/fixtures/SOURCES.md`). Подмена кладёт
сырое в приёмник так же, как настоящий `http.fetch` (`sink.keep` до
разбора), — проверяется порядок «сырое до разбора» и его задокументированные
исключения (ФИО и контакты вырезаются ДО записи).
"""
from __future__ import annotations

import dataclasses
import functools
import hashlib
import io
import json
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest

from indicators import http, issuer, sources
from indicators.store import Store

# Такт (ops/run.sh, TACT_TESTS): слой индикаторов на фикстурах и двойниках —
# быстрый, в сеть не ходит (тесты `network` такт исключает выражением).
pytestmark = pytest.mark.tact

FIXTURES = Path(__file__).parent / "fixtures"


def _response(url: str, body: bytes) -> http.Response:
    return http.Response(url=url, status=200, body=body,
                         fetched_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
                         headers={})


class FakeWeb:
    """Подмена `sources.fetch`: адрес → байты (или исключение), журнал запросов."""

    def __init__(self, routes):
        self.routes = routes
        self.calls: list[str] = []

    def __call__(self, url, *, sink=None, name=None, **_):
        self.calls.append(url)
        for pattern, answer in self.routes:
            if re.search(pattern, url):
                if isinstance(answer, Exception):
                    raise answer
                body = answer() if callable(answer) else answer
                response = _response(url, body)
                if sink is not None and name:
                    sink.keep(name, response)
                return response
        raise http.FetchError(url, 404, "HTTP 404")


def fixture(*parts) -> bytes:
    return FIXTURES.joinpath(*parts).read_bytes()


def _raw_files(store: Store, source: str) -> list[Path]:
    return sorted(p for p in (store.raw / source).rglob("*")
                  if p.is_file() and not p.name.endswith(".meta.json"))


def _meta(path: Path) -> dict:
    return json.loads(path.with_name(path.name + ".meta.json").read_text(encoding="utf-8"))


# ------------------------------------------------------------------ рынок


def test_quotes_of_the_issuer_and_the_peers(tmp_path, monkeypatch):
    web = FakeWeb([(r"boards/TQBR/securities\.json", fixture("moex_quote", "quotes.json"))])
    monkeypatch.setattr(sources, "fetch", web)
    result = sources.collect_moex_quote(store=Store(tmp_path))
    assert result.series[issuer.PRICE_SERIES][0].value == pytest.approx(1759.5)
    assert set(result.series) == {"moex.price.LENT", "moex.price.MGNT", "moex.price.X5"}
    assert "securities=LENT,X5,MGNT" in web.calls[0]


def test_a_morning_before_the_first_trade_is_a_quiet_day_not_a_failure(tmp_path, monkeypatch,
                                                                         capsys):
    """Тревога 03.10.2026 (суббота): ISS в 07:00 МСК обнуляет таблицу, и до первой
    сделки дня `LAST` пуст при `NUMTRADES` = 0 у всех бумаг. Утренний такт в
    выходной приходил на такую таблицу, критический сборщик давал «разобрано 0
    рядов», юнит трижды повторялся и слал «упало». Это тихий день, а не сбой:
    статус «ок» с причиной, точек нет, такт идёт дальше (код 0)."""
    from indicators import collect

    web = FakeWeb([(r"boards/TQBR/securities\.json",
                    fixture("moex_quote", "quotes_no_trades.json"))])
    monkeypatch.setattr(sources, "fetch", web)
    store = Store(tmp_path)

    result = sources.collect_moex_quote(store=store)
    assert result.series == {} and "торгов сегодня ещё не было" in result.idle
    assert "marketdata.columns=SECID,LAST,NUMTRADES,TRADINGSTATUS,UPDATETIME" in web.calls[0]

    report = sources.run_collectors(["moex_quote"], store)
    assert report["moex_quote"].startswith("ок: торгов сегодня ещё не было"), report
    assert store.load(issuer.PRICE_SERIES) is None, "новой точки цены нет"
    assert _raw_files(store, "moex"), "ответ биржи лежит в сыром архиве"
    assert collect.cmd_collect(["moex_quote"], store, retry=False, pause=0.0) == 0
    assert "ОТКАЗ КРИТИЧЕСКОГО" not in capsys.readouterr().out


@pytest.mark.parametrize("label, columns, rows", [
    ("сделки были, а последней сделки нет",
     ["SECID", "LAST", "NUMTRADES", "TRADINGSTATUS", "UPDATETIME"],
     [["LENT", None, 12, "T", "10:00:09"], ["MGNT", None, 0, "N", "07:00:09"]]),
    ("строки эмитента нет",
     ["SECID", "LAST", "NUMTRADES", "TRADINGSTATUS", "UPDATETIME"],
     [["MGNT", None, 0, "N", "07:00:09"], ["X5", None, 0, "N", "07:00:09"]]),
    ("в ответе нет числа сделок",
     ["SECID", "LAST", "UPDATETIME"],
     [["LENT", None, "07:00:09"], ["MGNT", None, "07:00:09"], ["X5", None, "07:00:09"]]),
    ("пустая таблица", ["SECID", "LAST", "NUMTRADES", "TRADINGSTATUS", "UPDATETIME"], []),
])
def test_an_empty_quote_that_is_not_a_quiet_day_stays_a_failure(tmp_path, monkeypatch,
                                                                label, columns, rows):
    """Тихим днём называется только «пусто у всех при нуле сделок и строке
    эмитента на месте». Всё прочее — мусор в ответе критического источника."""
    body = json.dumps({"marketdata": {"columns": columns, "data": rows}}).encode()
    monkeypatch.setattr(sources, "fetch",
                        FakeWeb([(r"boards/TQBR/securities\.json", body)]))
    store = Store(tmp_path)
    assert sources.collect_moex_quote(store=store).idle == "", label
    report = sources.run_collectors(["moex_quote"], store)
    assert report["moex_quote"] == "ОШИБКА: ответ получен, но разобрано 0 рядов", (label, report)


def test_the_security_card_gives_listing_level_and_issue_size(tmp_path, monkeypatch):
    """Карточка LENT: третий уровень листинга (с 12.08.2026), 115 985 197 акций."""
    body = fixture("moex_security", "security_LENT.json")
    monkeypatch.setattr(sources, "fetch", FakeWeb([(r"securities/LENT\.json", body)]))
    store = Store(tmp_path)
    monkeypatch.setattr(sources, "UTC_TODAY", lambda: "2026-09-27")
    report = sources.run_collectors(["moex_security"], store)
    assert report["moex_security"].startswith("ок"), report
    assert store.load("moex.security.LENT.listlevel").latest().value == 3.0
    assert store.load("moex.security.LENT.issuesize").latest().value == 115_985_197.0
    assert sources.security_changes(store) == []
    # Раз в сутки: второй такт дня источник не трогает.
    assert sources.run_collectors(["moex_security"], store)["moex_security"].startswith(
        "пропущен")

    # На следующий день биржа вернула акцию во второй уровень — это смена.
    changed = json.loads(body)
    rows = changed["description"]["data"]
    for row in rows:
        if row[0] == "LISTLEVEL":
            row[2] = "2"
    monkeypatch.setattr(sources, "fetch", FakeWeb(
        [(r"securities/LENT\.json", json.dumps(changed).encode())]))
    monkeypatch.setattr(sources, "UTC_TODAY", lambda: "2026-09-28")
    sources.run_collectors(["moex_security"], store)
    changes = sources.security_changes(store)
    assert changes == [dict(series="moex.security.LENT.listlevel", field="LISTLEVEL",
                            before=3.0, after=2.0, since="2026-09-28",
                            previous_day="2026-09-27")]


def test_bonds_are_found_by_issuer_code_not_by_name(tmp_path, monkeypatch):
    """Облигации группы — по кодам эмитентов ISS {5997, 4867}: в обращении только
    четыре выпуска «О'КЕЙ»; у ООО «Лента» все погашены (`is_traded = 0`)."""
    assert issuer.BOND_EMITTER_IDS == {5997, 4867}
    web = FakeWeb([
        (r"q=7814148471", fixture("bonds", "bonds_search_7814148471.json")),
        (r"q=7826087713", fixture("bonds", "bonds_search_7826087713.json")),
        (r"boards/TQCB/securities\.json", fixture("bonds", "bonds_tqcb.json")),
    ])
    monkeypatch.setattr(sources, "fetch", web)
    result = sources.collect_bonds(store=Store(tmp_path))
    prices = {k: v[0] for k, v in result.series.items() if k.endswith(".price")}
    assert set(prices) == {issuer.series(f"bond.{s}.price") for s in
                           ("RU000A1009Z8", "RU000A102BK7", "RU000A1084E6", "RU000A10CCN3")}
    assert prices[issuer.series("bond.RU000A10CCN3.price")].value == pytest.approx(1.0024)
    assert "оферта 2027-02-09" in prices[issuer.series("bond.RU000A10CCN3.price")].note
    assert result.series[issuer.series("bond.RU000A1084E6.yield")][0].value == \
        pytest.approx(0.1899)
    assert result.series[issuer.series("bonds.traded")][0].value == 4.0
    assert "securities=RU000A1009Z8,RU000A102BK7,RU000A1084E6,RU000A10CCN3" in web.calls[-1]
    assert len(web.calls) == 3, "два поиска по ИНН и одна доска — вежливость к ISS"


def test_the_key_rate_is_read_from_the_soap_answer(tmp_path, monkeypatch):
    monkeypatch.setattr(sources, "fetch", FakeWeb(
        [(r"DailyInfo\.asmx", fixture("cbr_key_rate", "key_rate.xml"))]))
    result = sources.collect_cbr_key_rate(since=date(2025, 9, 1))
    points = result.series["cbr.key_rate"]
    assert points and all(0.05 < p.value < 0.3 for p in points)
    assert all(re.fullmatch(r"\d{4}-\d{2}-\d{2}", p.period) for p in points)


# ------------------------------------------------------ лента раскрытия


def _feed_page(drop: int = 0) -> bytes:
    """Страница ленты из фикстуры; `drop` — снять N новейших сообщений (вчерашний снимок)."""
    page = fixture("lenta_disclosure", "regulatory_filings.html").decode("utf-8")
    app = sources.app_json(page)
    facts = app["components"]["ipjsc-lenta"]["material-facts"]
    app["components"]["ipjsc-lenta"]["material-facts"] = facts[drop:]
    start = page.index("App = ") + len("App = ")
    _, end = json.JSONDecoder().raw_decode(page[start:])
    return (page[:start] + json.dumps(app, ensure_ascii=False) + page[start + end:]).encode()


def _disclosure_web(feed: bytes) -> FakeWeb:
    routes = [(r"regulatory-filings/$", feed)]
    for fid in ("29130", "29047", "29046", "28466", "25850"):
        routes.append((rf"statement-of-material-facts/{fid}/$",
                       fixture("lenta_disclosure", f"fact_{fid}.html")))
    return FakeWeb(routes)


def _stamped(web: FakeWeb, stamp: str):
    """Подмена `fetch` с заданным моментом получения ответа: такты разных дней
    (и разные попытки одного тела) отличаются им, как на сервере."""
    def fetch(url, **kwargs):
        return dataclasses.replace(web(url, **kwargs), fetched_at=stamp)
    return fetch


def _disclosure_tact(monkeypatch, store: Store, day: str, web: FakeWeb, **kwargs) -> int:
    """Такт сбора ленты раскрытия «в день `day`» штатным `collect.cmd_collect`."""
    from indicators import collect

    monkeypatch.setattr(sources, "fetch", _stamped(web, f"{day}T14:25:00+00:00"))
    monkeypatch.setattr(sources, "UTC_TODAY", lambda: day)
    monkeypatch.setitem(sources.COLLECTORS, "lenta_disclosure", (functools.partial(
        sources.collect_lenta_disclosure, today=date.fromisoformat(day)),)
        + _DISCLOSURE_ENTRY[1:])
    return collect.cmd_collect(["lenta_disclosure"], store, **kwargs)


_DISCLOSURE_ENTRY = sources.COLLECTORS["lenta_disclosure"]


def _feed_until(day: str) -> bytes:
    """Страница ленты из фикстуры, какой она была на конец дня `day` (по Москве)."""
    page = fixture("lenta_disclosure", "regulatory_filings.html").decode("utf-8")
    app = sources.app_json(page)
    facts = app["components"]["ipjsc-lenta"]["material-facts"]
    app["components"]["ipjsc-lenta"]["material-facts"] = [
        f for f in facts if (sources.msk_day(f.get("date")) or "9999") <= day]
    start = page.index("App = ") + len("App = ")
    _, end = json.JSONDecoder().raw_decode(page[start:])
    return (page[:start] + json.dumps(app, ensure_ascii=False) + page[start + end:]).encode()


def test_disclosure_novelty_is_by_ids_and_the_event_date_by_item_1_7(tmp_path, monkeypatch):
    store = Store(tmp_path)
    # Вчера: снимок без двух новейших сообщений — первый снимок, новых 0.
    monkeypatch.setattr(sources, "fetch", _disclosure_web(_feed_page(drop=2)))
    first = sources.collect_lenta_disclosure(store=store, today=date(2026, 9, 10))
    new = first.series[issuer.series("disclosure.new")][0]
    assert new.value == 0.0 and "первый снимок" in new.note

    # Сегодня: два новых id — 29130 (совет директоров) и 29047 (листинг).
    web = _disclosure_web(_feed_page())
    monkeypatch.setattr(sources, "fetch", web)
    second = sources.collect_lenta_disclosure(store=store, today=date(2026, 9, 28))
    new = second.series[issuer.series("disclosure.new")][0]
    assert new.value == 2.0 and new.period == "2026-09-28"
    # Тела новых сообщений читаются первыми.
    bodies = [u for u in web.calls if "statement-of-material-facts" in u]
    assert bodies[:2] == [issuer.IR_SITE + "/ru/investors/regulatory-filings/ipjsc-lenta/"
                          "statement-of-material-facts/29130/",
                          issuer.IR_SITE + "/ru/investors/regulatory-filings/ipjsc-lenta/"
                          "statement-of-material-facts/29047/"]
    assert len(bodies) <= sources.MAX_EVENT_BODIES

    events = {p.period: p for p in second.series[issuer.series("disclosure.events")]}
    # Заседание совета проведено 16.09 (п. 1.7), опубликовано 17.09.
    assert "board:" in events["2026-09-16"].note
    assert "listing:" in events["2026-08-03"].note
    # Ряды классов — по дню ПУБЛИКАЦИИ из всей ленты.
    board = {p.period: p.value for p in second.series[issuer.series("disclosure.board")]}
    assert board["2026-09-17"] == 1.0
    listing = {p.period for p in second.series[issuer.series("disclosure.listing")]}
    assert "2026-08-04" in listing
    latest = second.series[issuer.series("disclosure.latest")][0]
    assert latest.period == "2026-09-17"


def test_the_signer_and_the_contacts_never_reach_the_disk(tmp_path, monkeypatch):
    """Блок подписанта и контакты пресс-релиза вырезаются ДО записи (D15).

    Фикстуры несут заглушки на месте ФИО и почт; на диске их быть не должно,
    а sha256 в мете — от очищенных байтов.
    """
    store = Store(tmp_path)
    web = _disclosure_web(_feed_page())
    monkeypatch.setattr(sources, "fetch", web)
    monkeypatch.setattr(sources, "MAX_EVENT_BODIES", 20)
    monkeypatch.setattr(sources, "BODY_BACKFILL_DAYS", 1000)   # «Молния» — 06.2025
    sources.collect_lenta_disclosure(store=store, today=date(2026, 9, 28))
    saved = {p.name: p for p in _raw_files(store, "lenta_disclosure")}
    assert {"event_29130.json", "event_29046.json", "event_25850.json"} <= set(saved)
    for name, path in saved.items():
        body = Store.read_raw(path)
        text = body.decode("utf-8")
        for needle in ("ПОДПИСАНТ-ЗАГЛУШКА", "КОНТАКТ-ЗАГЛУШКА", "почта-заглушка",
                       "3. Подпись", "Генеральный Директор"):
            assert needle not in text, (name, needle)
        assert _meta(path)["sha256"] == hashlib.sha256(body).hexdigest(), name
    press = json.loads(Store.read_raw(saved["event_29046.json"]))
    assert "[контакты вырезаны]" in press["text"]
    assert press["event_date"] == "2026-08-04" and press["kind"] == "listing"
    assert _meta(saved["event_29046.json"])["cleaned"] is True


def test_material_facts_are_refined_by_their_text():
    """«Существенное влияние» — это пресс-релизы: класс уточняется по тексту."""
    expected = {"29130": ("board", "2026-09-16"), "29047": ("listing", "2026-08-03"),
                "29046": ("listing", "2026-08-04"), "28466": ("results", "2025-10-31"),
                "25850": ("ma", "2025-06-26")}
    for fid, (kind, day) in expected.items():
        page = fixture("lenta_disclosure", f"fact_{fid}.html").decode("utf-8")
        detail = sources._component(sources.app_json(page), ("regulatory-filings-detail",))
        text = sources.clean_fact_text(detail["detail"])
        assert sources.refine_kind(sources.classify_disclosure(detail["name"]), text) == kind, fid
        assert sources.event_date(text) == day, fid


def test_a_class_refined_by_the_text_reaches_its_series_the_alarm_and_the_banner(
        tmp_path, monkeypatch, capsys):
    """T03 внешнего аудита: «существенное влияние» о покупке «Молнии» (26.06.2025)
    уточняется по телу в `ma` — и именно `ma` видят ряд класса, строка «СОБЫТИЕ»
    такта и плашка витрины (`payload._recent_events`). Ряд класса по названию при
    этом не убывает: сообщение остаётся и в `material`."""
    import model.payload as payload

    store = Store(tmp_path)
    monkeypatch.setattr(sources, "BODY_BACKFILL_DAYS", 30)
    assert _disclosure_tact(monkeypatch, store, "2025-06-25",
                            _disclosure_web(_feed_until("2025-06-25"))) == 0
    assert "СОБЫТИЕ" not in capsys.readouterr().err
    assert _disclosure_tact(monkeypatch, store, "2025-06-26",
                            _disclosure_web(_feed_until("2025-06-26"))) == 0
    err = capsys.readouterr().err
    ma = store.load(issuer.series("disclosure.ma"))
    assert ma is not None and "2025-06-26" in {p.period for p in ma.points}
    material = store.load(issuer.series("disclosure.material"))
    assert "2025-06-26" in {p.period for p in material.points}
    assert "СОБЫТИЕ: в ленте раскрытия — ma (публикация 2025-06-26)" in err

    class Frozen(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2025, 6, 27, 12, tzinfo=timezone.utc)

    monkeypatch.setattr(payload, "datetime", Frozen)
    events = payload._recent_events(store)
    assert [(e["date"], e["kind"], e["label"]) for e in events] == [
        ("2025-06-26", "ma", "сделки M&A")]


def test_a_class_learnt_a_day_late_still_raises_the_alarm_but_old_history_does_not(
        tmp_path, monkeypatch, capsys):
    """Тело прочитано на следующий день после публикации (вечерняя публикация,
    сбой загрузки, очередь) — «СОБЫТИЕ» даёт такт, который узнал класс: признак —
    день получения точки, а не «опубликовано сегодня». Историю, дочитанную через
    месяц, строка не поднимает; плашка окна 30 дней её ещё показывает."""
    store = Store(tmp_path)
    monkeypatch.setattr(sources, "BODY_BACKFILL_DAYS", 60)
    monkeypatch.setattr(sources, "MAX_EVENT_BODIES", 0)
    _disclosure_tact(monkeypatch, store, "2025-06-26", _disclosure_web(_feed_until("2025-06-26")))
    assert "СОБЫТИЕ" not in capsys.readouterr().err, "класс по названию — `material`"
    assert store.load(issuer.series("disclosure.ma")) is None
    monkeypatch.setattr(sources, "MAX_EVENT_BODIES", 3)
    _disclosure_tact(monkeypatch, store, "2025-06-27", _disclosure_web(_feed_until("2025-06-26")))
    assert "СОБЫТИЕ: в ленте раскрытия — ma (публикация 2025-06-26)" in capsys.readouterr().err

    late = Store(tmp_path / "late")
    monkeypatch.setattr(sources, "MAX_EVENT_BODIES", 0)
    _disclosure_tact(monkeypatch, late, "2025-06-26", _disclosure_web(_feed_until("2025-06-26")))
    monkeypatch.setattr(sources, "MAX_EVENT_BODIES", 3)
    capsys.readouterr()
    _disclosure_tact(monkeypatch, late, "2025-07-20", _disclosure_web(_feed_until("2025-06-26")))
    assert "СОБЫТИЕ" not in capsys.readouterr().err
    assert "2025-06-26" in {p.period for p in late.load(issuer.series("disclosure.ma")).points}


def test_the_press_release_about_a_purchase_is_a_deal():
    """«…ПОКУПАЕТ КОНТРОЛЬНУЮ ДОЛЮ…» — сделка: корень «покупк» это слово не ловил,
    и пресс-релиз уточнялся бы по слову «выручка» в результаты."""
    text = ("2. Содержание сообщения Пресс-релиз «ЛЕНТА» ВЫХОДИТ НА ДАЛЬНИЙ ВОСТОК – ПОКУПАЕТ "
            "КОНТРОЛЬНУЮ ДОЛЮ В РОЗНИЧНОЙ СЕТИ Выручка сети за год выросла")
    assert sources.refine_kind("material", text) == "ma"
    assert sources.refine_kind("board", text) == "board"


def test_an_unparsed_body_is_named_as_a_failure_and_read_again(tmp_path, monkeypatch, capsys):
    """T08 внешнего аудита: ответ 200 без текста сообщения не делает id
    «прочитанным». Источник получает отказ (собранное сохранено), следующий такт
    перечитывает тело — и событие появляется."""
    store = Store(tmp_path)
    web = _disclosure_web(_feed_page())
    web.routes.insert(1, (r"statement-of-material-facts/29130/$",
                          b"<html><body>tech works</body></html>"))
    code = _disclosure_tact(monkeypatch, store, "2026-09-28", web)
    captured = capsys.readouterr()
    assert code == 0, "лента восполнима: отказ второстепенного источника такт не роняет"
    assert "lenta_disclosure" in captured.err and "ТРЕВОГА" in captured.err
    report = json.loads((store.root / "collector_report.json").read_text(encoding="utf-8"))
    assert "29130 (попытка 1 из 3)" in report["sources"]["lenta_disclosure"]["status"]
    assert "29130" not in sources.saved_event_ids(store)
    assert sources.unparsed_attempts(store) == {"29130": (1, "2026-09-28")}
    names = {p.name for p in _raw_files(store, "lenta_disclosure")}
    assert "event_29130.unparsed.json" in names and "event_29130.json" not in names
    assert store.load(issuer.series("disclosure.board")) is not None, "собранное сохранено"

    # Сайт починился: следующий такт перечитывает тело, отказ снят.
    web = _disclosure_web(_feed_page())
    assert _disclosure_tact(monkeypatch, store, "2026-09-29", web) == 0
    assert any("/29130/" in u for u in web.calls)
    events = {p.period: p for p in store.load(issuer.series("disclosure.events")).points}
    assert "board:" in events["2026-09-16"].note
    assert "29130" in sources.saved_event_ids(store)
    report = json.loads((store.root / "collector_report.json").read_text(encoding="utf-8"))
    assert "lenta_disclosure" not in report["sources"]


def test_a_body_that_never_parses_is_tried_thrice_then_weekly_and_quietly(tmp_path, monkeypatch):
    """Тело, которое не читается никогда: три попытки подряд — с отказом источника,
    дальше — раз в неделю и молча (починенный разбор подхватит его сам, а запросы и
    внимание владельца не тратятся)."""
    store = Store(tmp_path)
    asked, failed = {}, {}
    start = date(2026, 10, 1)
    for offset in range(0, 19):
        day = (start + timedelta(days=offset)).isoformat()
        web = _disclosure_web(_feed_page())
        web.routes.insert(1, (r"statement-of-material-facts/29130/$",
                              b"<html><body>tech works</body></html>"))
        monkeypatch.setattr(sources, "fetch", _stamped(web, f"{day}T04:20:00+00:00"))
        result = sources.collect_lenta_disclosure(store=store, today=date.fromisoformat(day))
        asked[offset] = sum("/29130/" in u for u in web.calls)
        failed[offset] = bool(result.error)
    assert [d for d, n in asked.items() if n] == [0, 1, 2, 9, 16]
    assert [d for d, bad in failed.items() if bad] == [0, 1, 2]
    assert sources.unparsed_attempts(store)["29130"] == (5, "2026-10-17")
    # Две попытки одного дня (утренний и вечерний такт) — две копии, обе в счёте.
    twice = Store(tmp_path / "twice")
    for hour in ("04", "14"):
        web = _disclosure_web(_feed_page())
        web.routes.insert(1, (r"statement-of-material-facts/29130/$", b"<html></html>"))
        monkeypatch.setattr(sources, "fetch", _stamped(web, f"2026-10-01T{hour}:20:00+00:00"))
        sources.collect_lenta_disclosure(store=twice, today=start)
    assert sources.unparsed_attempts(twice)["29130"] == (2, "2026-10-01")


def test_a_compressed_same_day_copy_of_the_feed_is_the_latest(tmp_path, monkeypatch):
    """T09 внешнего аудита: боевая лента больше порога сжатия, и вторая копия дня
    ложится как `feed.json.1.gz` — прежний шаблон её не узнавал, последним снимком
    оставался первый, и уже учтённые id снова шли «новыми». Фикстура — 6 КБ:
    порог сжатия опускается, чтобы тест шёл по боевой ветке имён."""
    from indicators import store as store_module

    monkeypatch.setattr(store_module, "COMPRESS_ABOVE_BYTES", 1024)
    monkeypatch.setattr(sources, "MAX_EVENT_BODIES", 0)
    store = Store(tmp_path)
    day = date(2026, 10, 1)

    def run(drop: int) -> float:
        monkeypatch.setattr(sources, "fetch", _disclosure_web(_feed_page(drop=drop)))
        result = sources.collect_lenta_disclosure(store=store, today=day)
        return result.series[issuer.series("disclosure.new")][0].value

    assert run(2) == 0.0            # первый снимок — feed.json.gz
    assert run(1) == 1.0            # новое сообщение — feed.json.1.gz
    names = {p.name for p in _raw_files(store, "lenta_disclosure") if p.name.startswith("feed")}
    assert names == {"feed.json.gz", "feed.json.1.gz"}
    assert run(1) == 0.0, "тот же ответ ещё раз: новых нет"
    before, _ = sources.previous_feed(store)
    assert len(before) == len(sources.feed_items(sources.app_json(_feed_page(drop=1).decode())))
    assert run(0) == 1.0            # ещё одно — feed.json.2.gz, счёт — от .1.gz
    assert run(0) == 0.0
    for name in ("feed.json", "feed.json.gz", "feed.1.json", "feed.2.json.gz", "feed.json.3.gz"):
        assert sources.FEED_FILE.match(name), name
    for name in ("feed.json.meta.json", "feed.json.1.gz.meta.json", "feedback.json"):
        assert not sources.FEED_FILE.match(name), name


def test_the_feed_is_stored_only_when_it_changes(tmp_path, monkeypatch):
    """Страница — 1 МБ вёрстки; в архив идёт список ленты, и только новый."""
    store = Store(tmp_path)
    monkeypatch.setattr(sources, "MAX_EVENT_BODIES", 0)
    monkeypatch.setattr(sources, "fetch", _disclosure_web(_feed_page()))
    for _ in range(3):
        sources.collect_lenta_disclosure(store=store, today=date(2026, 9, 28))
    feeds = [p for p in _raw_files(store, "lenta_disclosure") if p.name.startswith("feed")]
    assert len(feeds) == 1
    items = json.loads(Store.read_raw(feeds[0]))
    assert items[-1]["id"] == "29130" and all("kind" in x for x in items)


def test_the_title_classes_of_the_mkpao_feed():
    cases = {
        "Проведение заседания совета директоров (наблюдательного совета) и его повестка дня":
            "board",
        "Раскрытие эмитентом консолидированной финансовой отчетности": "results",
        "Совершение подконтрольной эмитенту организацией существенной сделки": "ma",
        "Появление у эмитента подконтрольной организации, имеющей для него существенное значение":
            "ma",
        "Приобретение (наступление оснований для приобретения) размещенных эмитентом "
        "облигаций": "offer",
        "Перевод эмиссионных ценных бумаг эмитента из одного котировального списка в другой":
            "listing",
        "Выплаченные доходы по эмиссионным ценным бумагам эмитента": "coupon",
        "Дата, на которую определяются лица, имеющие право на осуществление прав": "record_date",
        "Сведения, направленные за пределы РФ для их раскрытия иностранным инвесторам": "other",
    }
    for title, kind in cases.items():
        assert sources.classify_disclosure(title) == kind, title


# ------------------------------------------------------------- датабук


def _databook_xlsx(extract: dict, *, scale: float = 1.0) -> bytes:
    """xlsx раскладки листа «Financials quarterly» из выписки фикстуры.

    Шапка «IAS 17» и метки кварталов, за ними блок «IFRS 16» (сборщик его не
    читает), ниже — раздел структуры SG&A с теми же метками и другими числами:
    сборщик берёт ПЕРВОЕ вхождение.
    """
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Content"
    sheet = book.create_sheet("Financials quarterly")
    sheet.cell(2, 1, "back to content")
    sheet.cell(4, 1, "Key Financial Results")
    sheet.cell(5, 1, "(in RUB mln)")
    quarters = extract["quarters"]
    ifrs_col = 2 + len(quarters) + 1
    sheet.cell(7, 2, "IAS 17")
    sheet.cell(7, ifrs_col, "IFRS 16")
    for j, label in enumerate(quarters):
        sheet.cell(8, 2 + j, label)
        sheet.cell(8, ifrs_col + j, label)
    row = 10
    for item in extract["rows"]:
        sheet.cell(row, 1, item["label"])
        for j, value in enumerate(item["values"]):
            if value is not None:
                sheet.cell(row, 2 + j, value * scale)
                sheet.cell(row, ifrs_col + j, value * 10)   # чужой блок — другие числа
        row += 1
    row += 2
    sheet.cell(row, 1, "SG&A Structure")
    for item in extract["rows"]:
        row += 1
        sheet.cell(row, 1, item["label"])
        for j in range(len(quarters)):
            sheet.cell(row, 2 + j, -1.0)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


EXTRACT = json.loads(fixture("lenta_databook", "financials_quarterly.json"))
DATABOOK_LINK = r"/upload/iblock/.*\.xlsx$"


def test_the_databook_quarterly_block_is_read_in_billions():
    blocks = sources.parse_databook(_databook_xlsx(EXTRACT))
    revenue = blocks["revenue"]
    assert sorted(revenue) == ["2025Q1", "2025Q2", "2025Q3", "2025Q4", "2026Q1", "2026Q2"]
    assert revenue["2026Q2"] == pytest.approx(341.593344, abs=1e-6)
    assert blocks["ebitda_pre16"]["2026Q2"] == pytest.approx(22.863, abs=1e-3)
    assert blocks["net_interest_pre16"]["2026Q2"] == pytest.approx(3.325, abs=1e-3), (
        "проценты — расходом со знаком «плюс», как их читает ретро-проверка")
    assert blocks["ebitda_margin_pre16"]["2026Q2"] == pytest.approx(22.863 / 341.593, abs=1e-5)
    assert blocks["revenue.remi"]["2026Q2"] > 0
    assert blocks["revenue.remi"]["2025Q1"] == 0.0, "до покупки в датабуке стоит ноль"
    assert all(v < -1.0 for v in blocks["payroll"].values()), (
        "расходы — со знаком отчётности (минус) и из P&L, а не из раздела SG&A (−1)")
    holed = json.loads(json.dumps(EXTRACT))
    holed["rows"][0]["values"][0] = None
    assert "2025Q1" not in sources.parse_databook(_databook_xlsx(holed))["revenue"], (
        "пустая ячейка — нет числа, а не ноль")


def test_a_broken_databook_is_an_error_not_an_empty_answer():
    import openpyxl

    with pytest.raises(sources.DatabookLayoutError, match="xlsx"):
        sources.parse_databook(b"<html>maintenance</html>")
    book = openpyxl.Workbook()
    book.active.title = "Other"
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(sources.DatabookLayoutError, match="Financials quarterly"):
        sources.parse_databook(buffer.getvalue())


def test_every_new_databook_version_is_kept_with_its_sha256(tmp_path, monkeypatch):
    """Защищённый источник: новая версия — в архив целиком, известная — не второй раз.

    Файл качается при новой ссылке или раз в неделю на прежней (замена без
    смены ссылки); пересчитанная история ложится рядом — винтажом.
    """
    store = Store(tmp_path)
    xlsx = _databook_xlsx(EXTRACT)
    web = FakeWeb([(r"/publications/$", fixture("lenta_databook", "publications.html")),
                   (DATABOOK_LINK, xlsx)])
    monkeypatch.setattr(sources, "fetch", web)
    day = date(2026, 9, 28)

    first = sources.collect_lenta_databook(store=store, today=day)
    assert "новая версия" in first.note
    published = first.series[issuer.series("databook.published")][0]
    assert published.period == "2026-08-03" and "Q22026_LENTA_DATABOOK.xlsx" in published.note
    assert first.series[issuer.series("databook.new_version")][0].period == day.isoformat()
    versions = sources.databook_versions(store)
    assert [v["sha256"] for v in versions] == [hashlib.sha256(xlsx).hexdigest()]
    stored = _raw_files(store, "lenta_databook")
    assert len(stored) == 1 and Store.read_raw(stored[0]) == xlsx
    for key, points in first.series.items():
        store.upsert(key, points)

    # Тот же день, та же ссылка — файл не качается.
    calls = len(web.calls)
    again = sources.collect_lenta_databook(store=store, today=day)
    assert issuer.series("databook.new_version") not in again.series
    assert all(not re.search(DATABOOK_LINK, u) for u in web.calls[calls:])

    # Через неделю — перепроверка: версия та же, второй копии нет.
    later = sources.collect_lenta_databook(store=store, today=day + timedelta(days=8))
    assert "не изменилась" in later.note
    assert len([v for v in sources.databook_versions(store) if v["kind"] == "version"]) == 1
    # И ещё через день перепроверка не повторяется (отметка перепроверки).
    calls = len(web.calls)
    sources.collect_lenta_databook(store=store, today=day + timedelta(days=9))
    assert all(not re.search(DATABOOK_LINK, u) for u in web.calls[calls:])

    # Новая ссылка (следующий релиз) с пересчитанной историей — новая версия и винтаж.
    page = fixture("lenta_databook", "publications.html").decode("utf-8").replace(
        "Q22026_LENTA_DATABOOK.xlsx", "Q32026_LENTA_DATABOOK.xlsx")
    revised = _databook_xlsx(EXTRACT, scale=1.01)
    monkeypatch.setattr(sources, "fetch", FakeWeb(
        [(r"/publications/$", page.encode()), (DATABOOK_LINK, revised)]))
    third = sources.collect_lenta_databook(store=store, today=day + timedelta(days=30))
    assert issuer.series("databook.new_version") in third.series
    for key, points in third.series.items():
        store.upsert(key, points)
    series = store.load(sources.databook_series_id("revenue"))
    assert len([p for p in series.points if p.period == "2026Q2"]) == 2, (
        "пересчитанная история не стирает прежнюю версию")
    assert len([v for v in sources.databook_versions(store) if v["kind"] == "version"]) == 2


def _renamed_sheet(body: bytes, title: str) -> bytes:
    """Тот же датабук с переименованным листом — «сайт сменил раскладку»."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(body))
    book[sources.DATABOOK_SHEET].title = title
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _databook_tact(monkeypatch, store: Store, day: date, web: FakeWeb, **kwargs) -> int:
    """Такт сбора датабука «в день `day`» штатным `collect.cmd_collect`."""
    from indicators import collect

    monkeypatch.setattr(sources, "fetch", web)
    monkeypatch.setattr(sources, "UTC_TODAY", lambda: day.isoformat())
    monkeypatch.setitem(sources.COLLECTORS, "lenta_databook", (functools.partial(
        sources.collect_lenta_databook, today=day),) + _DATABOOK_ENTRY[1:])
    return collect.cmd_collect(["lenta_databook"], store, pause=0.0, **kwargs)


_DATABOOK_ENTRY = sources.COLLECTORS["lenta_databook"]


def test_a_databook_that_failed_to_parse_keeps_the_alarm_until_the_parser_is_fixed(
        tmp_path, monkeypatch, capsys):
    """T01 внешнего аудита. Версия ложится в архив ДО разбора, и прежде повторный
    вызов видел «ту же ссылку» и отвечал «ок»: вечерний такт с повтором гасил
    тревогу сам, а квартальные ряды и «ВЫШЕЛ ОТЧЁТ» не появлялись никогда. Теперь
    отказ держится в каждом такте и в повторе, день публикации при этом собран;
    починенный разбор применяется ближайшим тактом к АРХИВНОЙ копии — без
    скачивания, с исходным винтажом, одной версией и одним файлом."""
    from indicators import collect

    store = Store(tmp_path)
    new_title = "Financials Quarterly (IAS 17)"
    web = FakeWeb([(r"/publications/$", fixture("lenta_databook", "publications.html")),
                   (DATABOOK_LINK, _renamed_sheet(_databook_xlsx(EXTRACT), new_title))])
    day = date(2026, 10, 29)

    # Вечерний такт с повтором — первый, кто видит новый датабук.
    assert _databook_tact(monkeypatch, store, day, web, retry=True) == collect.ALARM_EXIT
    captured = capsys.readouterr()
    assert "НЕВОСПОЛНИМЫЙ ПРОПУСК: lenta_databook" in captured.err
    assert "ВЫШЕЛ ОТЧЁТ" not in captured.err
    report = json.loads((store.root / collect.COLLECTOR_REPORT_NAME).read_text(encoding="utf-8"))
    status = report["sources"]["lenta_databook"]
    assert status["irrecoverable"] and "DatabookLayoutError" in status["status"]
    assert sum(bool(re.search(DATABOOK_LINK, u)) for u in web.calls) == 1, (
        "повтор разбирает архивную копию, файл второй раз не качается")
    assert store.load(issuer.series("databook.published")) is not None, "день публикации собран"
    assert store.load(sources.databook_series_id("revenue")) is None
    assert sources.databook_applied(store) == set()
    first_fetch = sources.databook_versions(store)[0]["fetched_at"]

    # Утро следующего дня (без повтора) и ещё через неделю — тревога держится.
    for later in (1, 8):
        assert _databook_tact(monkeypatch, store, day + timedelta(days=later), web) == \
            collect.ALARM_EXIT
    capsys.readouterr()

    # Разбор починен: ближайший такт кладёт версию в ряды из архива.
    monkeypatch.setattr(sources, "DATABOOK_SHEET", new_title)
    calls = len(web.calls)
    fixed_day = day + timedelta(days=9)
    assert _databook_tact(monkeypatch, store, fixed_day, web) == 0
    assert "ВЫШЕЛ ОТЧЁТ" in capsys.readouterr().err
    assert all(not re.search(DATABOOK_LINK, u) for u in web.calls[calls:])
    revenue = store.load(sources.databook_series_id("revenue"))
    assert revenue is not None and {p.fetched_at for p in revenue.points} == {first_fetch}
    signal = store.load(issuer.series("databook.new_version"))
    assert [(p.period, p.value) for p in signal.points] == [(fixed_day.isoformat(), 1.0)]
    versions = [v for v in sources.databook_versions(store) if v["kind"] == "version"]
    assert len(versions) == 1 and {versions[0]["sha256"]} == sources.databook_applied(store)
    assert len([p for p in _raw_files(store, "lenta_databook")
                if p.name.startswith("databook_") and "recheck" not in p.name]) == 1
    report = json.loads((store.root / collect.COLLECTOR_REPORT_NAME).read_text(encoding="utf-8"))
    assert "lenta_databook" not in report["sources"]

    # Дальше — как у любой учтённой версии: без скачивания и без нового сигнала.
    again = sources.collect_lenta_databook(store=store, today=fixed_day + timedelta(days=1))
    assert "тот же" in again.note and issuer.series("databook.new_version") not in again.series


def test_a_maintenance_page_on_the_databook_link_is_not_a_version(tmp_path, monkeypatch):
    """По ссылке xlsx отдали страницу техработ с кодом 200: отказ, тело — рядом в
    архиве, версией оно не становится — следующий такт качает файл снова, а не
    через неделю."""
    store = Store(tmp_path)
    answers = iter([b"<html>maintenance</html>", _databook_xlsx(EXTRACT)])
    web = FakeWeb([(r"/publications/$", fixture("lenta_databook", "publications.html")),
                   (DATABOOK_LINK, lambda: next(answers))])
    monkeypatch.setattr(sources, "fetch", web)
    report = sources.run_collectors(["lenta_databook"], store)
    assert report["lenta_databook"].startswith("ОШИБКА: DatabookLayoutError: по ссылке датабука "
                                               "пришёл не xlsx"), report
    assert sources.databook_versions(store) == []
    assert [p.name for p in _raw_files(store, "lenta_databook")][0].startswith("rejected_")
    report = sources.run_collectors(["lenta_databook"], store)
    assert report["lenta_databook"].startswith("ок"), report
    assert store.load(sources.databook_series_id("revenue")) is not None
    assert len(sources.databook_applied(store)) == 1


def test_two_databook_versions_of_one_day_are_both_counted(tmp_path, monkeypatch):
    """Файл заменили в день выхода: обе версии ложатся в ряды и обе учтены —
    вторая точка сигнала того же дня не сливается с первой (иначе версия считалась
    бы неприменённой и разбиралась бы в каждом такте)."""
    store = Store(tmp_path)
    day = date(2026, 10, 29)
    pages = fixture("lenta_databook", "publications.html")
    for number, scale in enumerate((1.0, 1.01), start=1):
        page = pages.decode("utf-8").replace("Q22026_LENTA", f"Q32026_v{number}_LENTA").encode()
        web = FakeWeb([(r"/publications/$", page),
                       (DATABOOK_LINK, _databook_xlsx(EXTRACT, scale=scale))])
        assert _databook_tact(monkeypatch, store, day, web) == 0
    signal = store.load(issuer.series("databook.new_version"))
    assert sorted(p.value for p in signal.points) == [1.0, 2.0]
    assert len(sources.databook_applied(store)) == 2
    revenue = store.load(sources.databook_series_id("revenue"))
    assert len([p for p in revenue.points if p.period == "2026Q2"]) == 2


def test_a_databook_without_a_single_revenue_number_is_a_layout_error():
    """Строка Total Sales есть, а чисел в ней нет — та же ошибка раскладки: версия
    без точек выручки легла бы в ряды пустой."""
    empty = json.loads(json.dumps(EXTRACT))
    for row in empty["rows"]:
        if row["label"].strip().lower() == "total sales":
            row["values"] = [None] * len(row["values"])
    with pytest.raises(sources.DatabookLayoutError, match="Total Sales"):
        sources.parse_databook(_databook_xlsx(empty))


def test_the_databook_collector_is_irrecoverable_and_protected():
    from indicators.store import PROTECTED_SOURCES

    assert "lenta_databook" in sources.IRRECOVERABLE
    assert "lenta_databook" in PROTECTED_SOURCES


# ------------------------------------------------------ «Работа России»


def _with_personal_data(page: bytes) -> bytes:
    """Ответ API, как он приходит: с контактами и свободными текстами (заглушки).

    Заглушки собираются во время теста: в файлах репозитория почтоподобных
    строк нет (гигиена, DESIGN раздел 1).
    """
    data = json.loads(page)
    mail = "hr" + "@" + "example.invalid"
    for item in data["results"]["vacancies"]:
        vacancy = item["vacancy"]
        vacancy["contact_person"] = "ФАМИЛИЯ-ЗАГЛУШКА Имя Отчество"
        vacancy["contact_list"] = [{"contact_type": "Телефон", "contact_value": "+7 900 000-00-00"},
                                   {"contact_type": "Эл. почта", "contact_value": mail}]
        vacancy["duty"] = "Звоните +7 (900) 000-00-00, пишите " + mail
        vacancy["addresses"] = {"address": [{"location": "ул. Заглушкина, 1"}]}
        vacancy["company"]["email"] = mail
        vacancy["company"]["phone"] = "+7 900 000 00 00"
    data["request"] = {"api": "v1"}
    return json.dumps(data, ensure_ascii=False).encode("utf-8")


def _trudvsem_web(*, failing: str | None = None) -> FakeWeb:
    routes = []
    for inn in issuer.employer_inns():
        page = fixture("trudvsem_vacancies", f"vacancies__inn_{inn}_p0.json")
        answer = (http.FetchError("http://opendata.trudvsem.ru", 460, "HTTP 460")
                  if inn == failing else _with_personal_data(page))
        routes.append((rf"company/inn/{inn}\?offset=0", answer))
    return FakeWeb(routes)


def test_vacancies_are_cleaned_before_they_touch_the_disk(tmp_path, monkeypatch):
    """ФИО, телефоны, почты и свободные тексты не ложатся на диск никогда (D15)."""
    store = Store(tmp_path)
    web = _trudvsem_web()
    monkeypatch.setattr(sources, "fetch", web)
    report = sources.run_collectors(["trudvsem_vacancies"], store)
    assert report["trudvsem_vacancies"].startswith("ок"), report
    files = [p for p in _raw_files(store, "trudvsem_vacancies")
             if sources.VACANCY_FILE.match(p.name)]
    assert len(files) == len(issuer.employer_inns())
    for path in files:
        body = Store.read_raw(path)
        text = body.decode("utf-8")
        for needle in ("contact_person", "contact_list", "ФАМИЛИЯ-ЗАГЛУШКА", "example.invalid",
                       "000-00-00", "000 00 00", "Заглушкина", "duty", "addresses",
                       '"request"', '"email"', '"phone"'):
            assert needle not in text, (path.name, needle)
        meta = _meta(path)
        assert meta["cleaned"] is True
        assert meta["sha256"] == hashlib.sha256(body).hexdigest(), (
            "sha256 — от очищенных байтов (задокументированное исключение)")
    totals = {inn: json.loads(fixture("trudvsem_vacancies", f"vacancies__inn_{inn}_p0.json"))
              ["meta"]["total"] for inn in issuer.employer_inns()}
    assert store.load(issuer.series("vacancies.total")).latest().value == sum(totals.values())
    for inn, total in totals.items():
        assert store.load(issuer.series(f"vacancies.total.{inn}")).latest().value == total
    assert all("offset=0&limit=100" in u for u in web.calls)
    # Раз в сутки: второй такт дня не ходит к медленному API.
    assert sources.run_collectors(["trudvsem_vacancies"], store)[
        "trudvsem_vacancies"].startswith("пропущен")


def test_a_failed_employer_keeps_the_rest_of_the_day(tmp_path, monkeypatch):
    """Отказ одного ИНН — «ОШИБКА» источника, но собранное по остальным сохранено."""
    store = Store(tmp_path)
    monkeypatch.setattr(sources, "fetch", _trudvsem_web(failing="7810495210"))
    report = sources.run_collectors(["trudvsem_vacancies"], store)
    assert report["trudvsem_vacancies"].startswith("ОШИБКА: не опрошены: 7810495210"), report
    assert store.load(issuer.series("vacancies.total.7814148471")) is not None
    assert store.load(issuer.series("vacancies.total.7810495210")) is None


def test_an_unreadable_vacancy_page_is_not_stored(tmp_path, monkeypatch):
    """Нечитаемый ответ «Работы России» на диск не кладётся: очистить его нельзя,
    а сырое с персональными данными не хранится (исключение D15). Это отказ."""
    store = Store(tmp_path)
    monkeypatch.setattr(sources, "fetch", FakeWeb([(r"trudvsem", b"<html>460</html>")]))
    report = sources.run_collectors(["trudvsem_vacancies"], store)
    assert report["trudvsem_vacancies"].startswith("ОШИБКА"), report
    assert not (store.raw / "trudvsem_vacancies").exists() or not _raw_files(
        store, "trudvsem_vacancies")


def test_clean_vacancy_keeps_only_the_white_list():
    vacancy = {"id": "x", "salary_min": 50000, "salary_max": 60000, "job-name": "кассир",
               "contact_person": "Кто-то", "contact_list": [{"a": 1}], "duty": "текст",
               "region": {"region_code": "7800000000000", "name": "Санкт-Петербург",
                          "extra": "x"},
               "company": {"inn": "7814148471", "name": "ООО \"ЛЕНТА\"", "hr-agency": "y"}}
    cleaned = sources.clean_vacancy(vacancy)
    assert set(cleaned) == {"id", "salary_min", "salary_max", "job-name", "region", "company"}
    assert cleaned["region"] == {"region_code": "7800000000000", "name": "Санкт-Петербург"}
    assert cleaned["company"] == {"inn": "7814148471", "name": "ООО \"ЛЕНТА\""}


# ------------------------------------------------------------- гигиена


EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE = sources.PHONE_RE
INITIALS = re.compile(r"\b[А-ЯЁ]\.\s?[А-ЯЁ]\.\s?[А-ЯЁ][а-яё]{2,}|\b[А-ЯЁ][а-яё]{2,}\s[А-ЯЁ]\.\s?[А-ЯЁ]\.")


def test_fixtures_carry_no_personal_data():
    """Ни почт, ни телефонов, ни полей контактов, ни «И.О. Фамилия» в фикстурах."""
    for path in sorted(FIXTURES.rglob("*")):
        if not path.is_file():
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        assert not EMAIL.search(text), (path.name, EMAIL.search(text).group(0))
        assert not PHONE.search(text), (path.name, PHONE.search(text).group(0))
        if path.suffix in (".json", ".html", ".xml"):   # описание фикстур называет поля
            assert "contact_person" not in text and "contact_list" not in text, path.name
        assert not INITIALS.search(text), (path.name, INITIALS.search(text).group(0))


def test_the_user_agent_names_the_project_and_not_a_person():
    assert http.USER_AGENT == ("tzi-850-lenta/1.0 "
                               "(+https://github.com/ML371KL/temp-zero-inode-850-lenta)")
    assert not EMAIL.search(http.USER_AGENT)
    http.USER_AGENT.encode("latin-1")   # заголовки HTTP — latin-1


def test_the_legal_entities_have_a_verification_note():
    """ИНН юрлиц сверены с карточкой эмитента ISS и/или ответом «Работы России»;
    сверки по ЕГРЮЛ нет — это отмечено, а не скрыто."""
    for entity in issuer.LEGAL_ENTITIES:
        assert re.fullmatch(r"\d{10}", entity["inn"]), entity
        assert entity["verified"] and entity["egrul_checked"] is False
    assert set(issuer.employer_inns()) == {"7814148471", "7826087713", "6674121179",
                                           "7810495210"}


# ------------------------------------------------------ после сбора (такт)


def test_the_tact_says_a_report_is_out_and_recounts_the_salary_index(tmp_path, monkeypatch,
                                                                     capsys):
    """После сбора: «ВЫШЕЛ ОТЧЁТ» при новой версии датабука, пересчёт индекса вилок
    после удачного обхода «Работы России», «СОБЫТИЕ» по классам тревоги."""
    from indicators import collect

    routes = [(r"/publications/$", fixture("lenta_databook", "publications.html")),
              (DATABOOK_LINK, _databook_xlsx(EXTRACT))]
    routes += _trudvsem_web().routes
    monkeypatch.setattr(sources, "fetch", FakeWeb(routes))
    store = Store(tmp_path)
    code = collect.cmd_collect(["lenta_databook", "trudvsem_vacancies"], store)
    captured = capsys.readouterr()
    assert code == 0, captured.err
    assert "ВЫШЕЛ ОТЧЁТ" in captured.err
    assert "зарплатный индекс: звеньев нет" in captured.out
    assert store.load(sources.databook_series_id("revenue")) is not None

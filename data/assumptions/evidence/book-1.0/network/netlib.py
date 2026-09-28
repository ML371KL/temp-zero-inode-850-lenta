# -*- coding: utf-8 -*-
"""Общие функции листа «Сеть и выручка» (книга Ленты 1.0).

Первичка — только через LENTA_PRIMARY_DIR (по умолчанию — reference/primary папки передачи,
вычисляется относительно этого файла). Росстат (ИПЦ) — файл рядом со скриптом
`rosstat_ipc_mes_07-2026.xlsx` (https://rosstat.gov.ru/storage/mediabank/ipc_mes_07-2026.xlsx,
sha256 ace8990f…), переопределяется LENTA_CPI_FILE.

Ничего не пишет, кроме того, что просят вызывающие скрипты.
"""
from __future__ import annotations

import hashlib
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
# book-draft/evidence/book-1.0/network → папка передачи = четыре уровня вверх
H_DIR = os.path.normpath(os.path.join(HERE, '..', '..', '..', '..'))
PRIMARY = os.environ.get('LENTA_PRIMARY_DIR', os.path.join(H_DIR, 'reference', 'primary'))
CPI_FILE = os.environ.get('LENTA_CPI_FILE', os.path.join(HERE, 'rosstat_ipc_mes_07-2026.xlsx'))

DATABOOK = 'Lenta_Q22026_DATABOOK.xlsx'
SHA = {  # из reference/primary/MANIFEST.md
    'Lenta_Q22026_DATABOOK.xlsx': '0d290ed28fc13a0b44cdf3dd9e0f5dd519b2fa92f9206f4b4d731a01f95b3339',
    'ifrs_mkpao/ifrs_1H2026.pdf': '235bdec9366875c0b1e7d410e4b2be15072ecde04d40fd56777947c5f4588797',
    'ifrs_mkpao/ifrs_FY2025.pdf': '4864a271babf7b34d17a4220f264bcc4d605f1e39964a65377c5b083c5de53a4',
    'Lenta_strategy-presentation-2028.pdf': 'a06c175fd9ed1971fbc0a6dc020f44afc65d5171200076d6063ad76b44ad4530',
}
CPI_SHA = 'ace8990fe8358173f743987a256eaef71501b06b5c4e5fe865b28046776ea412'


def sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def primary(rel: str, check: bool = True) -> str:
    p = os.path.join(PRIMARY, *rel.split('/'))
    if not os.path.exists(p):
        sys.exit(f'нет первички: {p} (задайте LENTA_PRIMARY_DIR)')
    if check and rel in SHA:
        got = sha256(p)
        if got != SHA[rel]:
            sys.exit(f'sha256 {rel}: {got} != MANIFEST {SHA[rel]}')
    return p


# ------------------------------------------------------------------ датабук
# Строки листа Operating Results (1-based) и ожидаемые подписи: макет проверяется, а не угадывается.
OPER_BLOCKS = {
    'revenue': (8, ['Total Sales', 'Retail Sales', 'Hypermarket', 'Supermarkets', 'Convenience stores', 'Utkonos',
                    'Drogerie', 'Remi', 'Dom Lenta', 'Other formats', 'Wholesales']),
    'avg_ticket': (22, ['Average Ticket', 'Hypermarket', 'Supermarkets', 'Convenience stores', 'Utkonos', 'Drogerie',
                        'Remi', 'Dom Lenta', 'Other formats']),
    'tickets': (32, ['Number of Tickets', 'Hypermarket', 'Supermarkets', 'Convenience stores', 'Utkonos', 'Drogerie',
                     'Remi', 'Dom Lenta', 'Other formats']),
    'lfl_group': (42, ['LFL Retail Sales', 'Retail traffic', 'Retail ticket']),
    'lfl_hyper': (47, ['LFL Sales, Hypermarket', 'Traffic', 'Ticket']),
    'lfl_super': (51, ['LFL Sales, Supermarkets', 'Traffic', 'Ticket']),
    'lfl_conv': (55, ['LFL Sales, Convenience', 'Traffic', 'Ticket']),
    'lfl_droge': (59, ['LFL Sales, Drogerie', 'Traffic', 'Ticket']),
    'stores': (63, ['Total Retail Stores', 'Hypermarkets', 'Supermarkets', 'Convenience stores', 'Drogerie', 'Remi',
                    'Dom Lenta', 'Other formats']),
    'net_open': (72, ['Total Net Store Openings', 'Hypermarkets', 'Supermarkets', 'Convenience stores', 'Drogerie',
                      'Remi', 'Dom Lenta', 'Other formats']),
    'area': (81, ['Total Selling Space', 'Hypermarkets', 'Supermarkets', 'Convenience stores', 'Drogerie', 'Remi',
                  'Dom Lenta', 'Other formats']),
    'new_area': (90, ['New Selling Space', 'Hypermarkets', 'Supermarkets', 'Convenience stores', 'Drogerie', 'Remi',
                      'Dom Lenta', 'Other formats']),
    'online': (117, ['Total Online Sales', 'Online Partners', 'Lenta Online', 'Utkonos']),
}
MONTH_BLOCKS = {
    'revenue': (9, ['Total Sales', 'Retail Sales', 'Hypermarket', 'Supermarkets', 'Convenience stores', 'Utkonos',
                    'Drogerie', 'Remi', 'Dom Lenta', 'Other formats', 'Wholesales']),
    'stores': (42, ['Total Retail Stores', 'Hypermarkets', 'Supermarkets', 'Convenience stores', 'Drogerie', 'Remi',
                    'Dom Lenta', 'Other formats']),
    'net_open': (51, ['Total Net Store Openings', 'Hypermarkets', 'Supermarkets', 'Convenience stores', 'Drogerie',
                      'Remi', 'Dom Lenta', 'Other formats']),
    'area': (60, ['Total Selling Space', 'Hypermarkets', 'Supermarkets', 'Convenience stores', 'Drogerie', 'Remi',
                  'Dom Lenta', 'Other formats']),
    'new_area': (69, ['New Selling Space', 'Hypermarkets', 'Supermarkets', 'Convenience stores', 'Drogerie', 'Remi',
                      'Dom Lenta', 'Other formats']),
}
# короткие имена строк датабука (форматы)
FMT = {'Total Sales': 'total', 'Retail Sales': 'retail', 'Hypermarket': 'hyper', 'Hypermarkets': 'hyper',
       'Supermarkets': 'super', 'Convenience stores': 'conv', 'Utkonos': 'utkonos', 'Drogerie': 'droge',
       'Remi': 'remi', 'Dom Lenta': 'diy', 'Other formats': 'other', 'Wholesales': 'wholesale',
       'Average Ticket': 'total', 'Number of Tickets': 'total', 'Total Retail Stores': 'total',
       'Total Net Store Openings': 'total', 'Total Selling Space': 'total', 'New Selling Space': 'total',
       'Total Online Sales': 'total', 'Online Partners': 'partners', 'Lenta Online': 'own', }
LFL_NAMES = ['sales', 'traffic', 'ticket']
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def load_databook():
    """→ dict(oper={block: {fmt: {period: value}}}, month={block: {fmt: {'YYYY-MM': value}}}).

    Периоды листа Operating Results: 'YYYYQn', 'YYYYFY', '2026H1'. Нули датабука в строках форматов до их
    появления — это «нет формата»; они сохраняются как 0.0, интерпретацию делают вызывающие.
    """
    import openpyxl
    wb = openpyxl.load_workbook(primary(DATABOOK), read_only=True, data_only=True)
    ws = wb['Operating Results']
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[6]
    cols = {}
    for c, v in enumerate(hdr):
        if not isinstance(v, str):
            continue
        s = v.replace(' ', '')
        m = re.fullmatch(r'([1-4])Q(\d{4})', s)
        if m:
            cols[c] = f'{m.group(2)}Q{m.group(1)}'
        m = re.fullmatch(r'FY(\d{4})', s)
        if m:
            cols[c] = f'{m.group(1)}FY'
        m = re.fullmatch(r'1H(\d{4})', s)
        if m:
            cols[c] = f'{m.group(1)}H1'
    oper = {}
    for block, (r0, labels) in OPER_BLOCKS.items():
        out = {}
        for i, lab in enumerate(labels):
            row = rows[r0 - 1 + i]
            got = str(row[0] or '').strip()
            if not got.startswith(lab):
                sys.exit(f'макет датабука изменился: строка {r0 + i} «{got}», ожидалось «{lab}…»')
            key = LFL_NAMES[i] if block.startswith('lfl_') else FMT[lab]
            out[key] = {p: _num(row[c]) for c, p in cols.items()}
        oper[block] = out
    ws = wb['Monthly Operating Results']
    rows = list(ws.iter_rows(values_only=True))
    yr, mo = rows[6], rows[7]
    mcols, cur = {}, None
    for c in range(len(mo)):
        if isinstance(yr[c], (int, float)):
            cur = int(yr[c])
        if isinstance(mo[c], str) and mo[c][:3] in MONTHS and cur:
            mcols[c] = f'{cur}-{MONTHS.index(mo[c][:3]) + 1:02d}'
    month = {}
    for block, (r0, labels) in MONTH_BLOCKS.items():
        out = {}
        for i, lab in enumerate(labels):
            row = rows[r0 - 1 + i]
            got = str(row[0] or '').strip()
            if not got.startswith(lab):
                sys.exit(f'макет месячного листа изменился: строка {r0 + i} «{got}», ожидалось «{lab}…»')
            out[FMT[lab]] = {p: _num(row[c]) for c, p in mcols.items()}
        month[block] = out
    return dict(oper=oper, month=month)


def halves_of(q: dict, y: int):
    """Сумма кварталов → полугодия года y: (H1, H2); None, если квартала нет."""
    def s(a, b):
        va, vb = q.get(f'{y}Q{a}'), q.get(f'{y}Q{b}')
        return None if va is None or vb is None else va + vb
    return s(1, 2), s(3, 4)


def half_key(y: int, h: int) -> str:
    return f'{y}H{h}'


def quarters_of_half(p: str):
    y, h = int(p[:4]), int(p[5])
    return [f'{y}Q{2 * h - 1}', f'{y}Q{2 * h}']


def eop_quarter(p: str) -> str:
    y, h = int(p[:4]), int(p[5])
    return f'{y}Q{2 * h}'


# ------------------------------------------------------------------ ИПЦ Росстата
def load_cpi():
    """→ {'all'|'food'|'nonfood': {(y, m): уровень}} по цепочке месячных индексов листов 01/02/03."""
    import openpyxl
    if not os.path.exists(CPI_FILE):
        sys.exit(f'нет файла ИПЦ: {CPI_FILE}')
    if sha256(CPI_FILE) != CPI_SHA:
        sys.exit('sha256 файла ИПЦ не совпадает с записанным')
    wb = openpyxl.load_workbook(CPI_FILE, read_only=True, data_only=True)
    out = {}
    for sh, name in (('01', 'all'), ('02', 'food'), ('03', 'nonfood')):
        rows = list(wb[sh].iter_rows(values_only=True))
        years = rows[3]
        mom = {}
        for mi in range(12):
            row = rows[5 + mi]
            for c, y in enumerate(years):
                if isinstance(y, (int, float)) and isinstance(row[c], (int, float)):
                    mom[(int(y), mi + 1)] = float(row[c])
        lev, cur = {}, 100.0
        for k in sorted(mom):
            cur *= mom[k] / 100.0
            lev[k] = cur
        out[name] = lev
    return out


def cpi_yoy(cpi: dict, name: str, y: int, months) -> float | None:
    """Рост среднего уровня за месяцы периода к тем же месяцам прошлого года, доля."""
    lev = cpi[name]
    try:
        a = sum(lev[(y, m)] for m in months) / len(months)
        b = sum(lev[(y - 1, m)] for m in months) / len(months)
    except KeyError:
        return None
    return a / b - 1


PERIOD_MONTHS = {'Q1': (1, 2, 3), 'Q2': (4, 5, 6), 'Q3': (7, 8, 9), 'Q4': (10, 11, 12),
                 'H1': (1, 2, 3, 4, 5, 6), 'H2': (7, 8, 9, 10, 11, 12), 'FY': tuple(range(1, 13))}


def cpi_period(cpi: dict, name: str, p: str) -> float | None:
    return cpi_yoy(cpi, name, int(p[:4]), PERIOD_MONTHS[p[4:]])


# ------------------------------------------------------------------ PDF и тексты первички
def pdf_text(rel: str, pages) -> str:
    from pypdf import PdfReader
    r = PdfReader(primary(rel))
    return '\n'.join(r.pages[i - 1].extract_text() or '' for i in pages)


def text_file(rel: str) -> str:
    with open(primary(rel, check=False), encoding='utf-8') as f:
        return f.read()


def rub_thousands(s: str) -> float:
    """«9 359 828» (тыс. руб.) → млрд руб."""
    return float(re.sub(r'\s', '', s)) / 1e6


# ------------------------------------------------------------------ правило ядра: эффективная площадь
def immature(cohorts, maturity, density=1.0, n_dense=0):
    """Незрелая часть когорт (старшая первой в списке): Σ n·(1 − m(age)·(d, если age < n_dense, иначе 1))."""
    return sum(n * (1 - maturity[min(a, len(maturity) - 1)] * (density if a < n_dense else 1.0))
               for a, n in enumerate(reversed(cohorts)))


def effective_area_step(cohorts, area0, eff_start, opened, area_end, maturity, closed_productivity,
                        density=1.0, n_dense=0):
    """Шаг «сеть» ядра 850oa (model/core.py; ops/tools/reanchor.effective_area_step) на фактах.

    Закрытия — из тождества area0 + opened − closed = area_end; когорта полугодия дозревает до d.
    Возвращает start/end/average по правилу, reset (уровень, который ядро восстановит на новом якоре), δ.
    """
    closed = area0 + opened - area_end
    maturing = sum(cohorts[-a] * (maturity[a] - maturity[a - 1]) * (density if a <= n_dense else 1.0)
                   for a in range(1, min(len(maturity), len(cohorts) + 1)))
    end = eff_start - closed * closed_productivity + opened * maturity[0] * density + maturing
    rotated = cohorts[1:] + [opened]
    n_dense_new = min(len(rotated), n_dense + 1) if density != 1.0 else n_dense
    reset = area_end - immature(rotated, maturity, density, n_dense_new)
    return dict(start=eff_start, end=end, average=(eff_start + end) / 2, reset=reset, delta=end - reset,
                closed=closed, cohorts=rotated, n_dense=n_dense_new)


def eff_area_chain(area_ends, gross, maturity, closed_productivity, density=1.0, n_dense_hist=0):
    """Непрерывный индекс эффективной площади правилом ядра, сведённый к уровню ядра на последнем конце.

    area_ends: [A(t0), A(t1), …, A(tn)] физическая площадь на концы полугодий;
    gross: [когорты до t0 (старшая первой, len = len(maturity) − 1)…] + [открытия (t0,t1] … (tn-1,tn]].
    Возвращает средние по полугодиям (t0,t1] … в конвенции ядра на якоре tn (вариант «в» листа Магнита).
    """
    k = len(maturity) - 1
    cohorts = list(gross[:k])
    opened_seq = list(gross[k:])
    assert len(opened_seq) == len(area_ends) - 1
    start = area_ends[0] - immature(cohorts, maturity, density, n_dense_hist)
    idx = [start]
    steps = []
    n_dense = n_dense_hist
    for i, op in enumerate(opened_seq):
        s = effective_area_step(cohorts, area_ends[i], idx[-1], op, area_ends[i + 1], maturity,
                                closed_productivity, density, n_dense)
        steps.append(s)
        # следующий шаг начинается с уровня «физическая − незрелая» нового якоря (как перезаякоривание),
        # индекс непрерывен: сдвиг δ переносится на всю предшествующую историю
        cohorts, n_dense = s['cohorts'], s['n_dense']
        idx.append(s['reset'])
    # непрерывный индекс: приросты по правилу, уровень — reset последнего якоря
    level_end = steps[-1]['reset']
    pts = [level_end]
    for s in reversed(steps):
        pts.append(pts[-1] - (s['end'] - s['start']))
    pts = list(reversed(pts))          # уровни на концах t0 … tn
    avgs = [(pts[i] + pts[i + 1]) / 2 for i in range(len(steps))]
    return dict(levels=pts, averages=avgs, steps=steps)


def fmt(x, nd=1):
    if x is None:
        return '—'
    s = f'{x:,.{nd}f}'.replace(',', ' ').replace('.', ',')
    return s

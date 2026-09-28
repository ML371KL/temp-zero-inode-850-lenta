"""Общие функции листа фактов книги 1.0 (Лента 850).

Первичка читается только через переменную LENTA_PRIMARY_DIR (по умолчанию —
reference/primary рядом с book-draft, путь выводится от места этого файла).
Абсолютных путей пользователя в коде нет.
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import json
import os
import re
from functools import lru_cache
from pathlib import Path

HERE = Path(__file__).resolve().parent
BOOK_DRAFT = HERE.parents[2]          # .../book-draft
HANDOFF = BOOK_DRAFT.parent           # .../lenta-850-handoff


def primary_dir() -> Path:
    return Path(os.environ.get("LENTA_PRIMARY_DIR", HANDOFF / "reference" / "primary"))


def research_dir() -> Path:
    return Path(os.environ.get("LENTA_RESEARCH_DIR", HANDOFF / "research"))


def facts_out_dir() -> Path:
    return Path(os.environ.get("LENTA_FACTS_OUT", BOOK_DRAFT / "facts"))


def magnit_dir() -> Path:
    # только для сверки строки аналогов (read-only); отсутствие не ломает сборку
    return Path(os.environ.get("MAGNIT_850OA_DIR", HANDOFF.parent / "magnit-850oa"))


# ----------------------------------------------------------------- хэши
def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


@lru_cache(maxsize=1)
def manifest() -> dict[str, dict]:
    """Строки MANIFEST.md: путь → {sha256, url, date, desc}."""
    out: dict[str, dict] = {}
    txt = (primary_dir() / "MANIFEST.md").read_text(encoding="utf-8")
    for line in txt.splitlines():
        if not line.startswith("|") or line.startswith("|---"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 5 or not re.fullmatch(r"[0-9a-f]{64}", cells[3] or ""):
            continue
        out[cells[0]] = {"url": cells[1], "date": cells[2], "sha256": cells[3], "desc": cells[4]}
    return out


# ----------------------------------------------------------------- датабук
_PERIOD_RE = [
    (re.compile(r"^FY\s*(\d{4})\s*restated$", re.I), lambda m: f"FY{m.group(1)}R"),
    (re.compile(r"^FY\s*(\d{4})$", re.I), lambda m: f"FY{m.group(1)}"),
    (re.compile(r"^1[HН]\s*(\d{4})$"), lambda m: f"{m.group(1)}H1"),   # латинская H и кириллическая Н
    (re.compile(r"^([1-4])Q\s*(\d{4})$"), lambda m: f"{m.group(2)}Q{m.group(1)}"),
]


def norm_period(label) -> str | None:
    if label is None:
        return None
    if isinstance(label, _dt.datetime):
        return label.strftime("%Y-%m-%d")
    if isinstance(label, int):
        return f"FY{label}"
    s = str(label).strip()
    for rx, fn in _PERIOD_RE:
        m = rx.match(s)
        if m:
            return fn(m)
    return None


def norm_label(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip() if s is not None else ""


class Sheet:
    """Лист датабука: столбец → (база, период); строка → подпись в столбце A.

    Базы определяются маркерами «IAS 17» / «IFRS 16» в строке basis_row.
    """

    def __init__(self, ws, header_row: int = 8, basis_row: int = 7, scale: float = 1e-6, name: str = ""):
        self.ws = ws
        self.name = name or ws.title
        self.scale = scale
        markers = []
        for c in ws[basis_row]:
            v = norm_label(c.value)
            if v in ("IAS 17", "IFRS 16"):
                markers.append((c.column, "ias17" if v == "IAS 17" else "ifrs16"))
        markers.sort()
        self.cols: dict[tuple[str, str], int] = {}
        for c in ws[header_row]:
            per = norm_period(c.value)
            if per is None:
                continue
            basis = None
            for col, b in markers:
                if col <= c.column:
                    basis = b
            if basis is None:
                continue
            self.cols.setdefault((basis, per), c.column)
        self.rows: dict[str, int] = {}
        for r in range(1, ws.max_row + 1):
            lab = norm_label(ws.cell(r, 1).value)
            if lab and lab not in self.rows:
                self.rows[lab] = r

    def row(self, label: str) -> int:
        lab = norm_label(label)
        if lab not in self.rows:
            raise KeyError(f"{self.name}: нет строки «{label}»")
        return self.rows[lab]

    def cell(self, label: str, basis: str, period: str):
        """Возвращает (значение в млрд или None, ссылка на ячейку)."""
        from openpyxl.utils import get_column_letter

        r = self.row(label)
        c = self.cols.get((basis, period))
        if c is None:
            return None, None
        v = self.ws.cell(r, c).value
        ref = f"{self.name}!{get_column_letter(c)}{r}"
        if isinstance(v, (int, float)):
            return v * self.scale, ref
        return None, ref

    def periods(self, basis: str) -> list[str]:
        return [p for (b, p) in self.cols if b == basis]


# ----------------------------------------------------------------- PDF
_NUM_RE = re.compile(r"(?<!\d)(\d{1,3}(?:(?:\\~ ?|[   ,])\d{3})+(?!\d)|\d+)")


@lru_cache(maxsize=None)
def pdf_pages(relpath: str) -> tuple[str, ...]:
    import pypdf

    r = pypdf.PdfReader(str(primary_dir() / relpath))
    return tuple((p.extract_text() or "") for p in r.pages)


@lru_cache(maxsize=None)
def text_file(relpath: str) -> str:
    return (primary_dir() / relpath).read_text(encoding="utf-8", errors="replace")


def numbers_in(text: str) -> set[int]:
    out = set()
    for m in _NUM_RE.finditer(text):
        d = re.sub(r"\D", "", m.group(1))
        if d:
            out.add(int(d))
    return out


@lru_cache(maxsize=None)
def pdf_page_numbers(relpath: str, page: int) -> frozenset:
    pages = pdf_pages(relpath)
    if page < 1 or page > len(pages):
        return frozenset()
    return frozenset(numbers_in(pages[page - 1]))


def find_number_pages(relpath: str, value: int, pages: range | None = None) -> list[int]:
    allp = pdf_pages(relpath)
    rng = pages or range(1, len(allp) + 1)
    return [p for p in rng if abs(int(value)) in pdf_page_numbers(relpath, p)]


def squash(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace(" ", " ")).strip()


def pdf_page_has_text(relpath: str, page: int, quote: str) -> bool:
    pages = pdf_pages(relpath)
    return 1 <= page <= len(pages) and squash(quote) in squash(pages[page - 1])


def r6(x):
    return None if x is None else round(float(x), 6)


def dump_json(obj, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    txt = json.dumps(obj, ensure_ascii=False, indent=1, sort_keys=False)
    path.write_text(txt + "\n", encoding="utf-8", newline="\n")

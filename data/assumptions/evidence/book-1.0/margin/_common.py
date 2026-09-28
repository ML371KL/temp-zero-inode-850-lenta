# -*- coding: utf-8 -*-
"""Общие функции листа «Маржа» книги Ленты 1.0.

Первичка читается из каталога LENTA_PRIMARY_DIR (по умолчанию — reference/primary
папки передачи, путь вычисляется от расположения этого файла, без литералов
C:\\Users). Никаких чисел книги здесь нет: только чтение первички.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import sys
from functools import lru_cache
from pathlib import Path

HERE = Path(__file__).resolve().parent
# book-draft/evidence/book-1.0/margin -> папка передачи H = parents[3]
H_DIR = HERE.parents[3]
PRIMARY = Path(os.environ.get("LENTA_PRIMARY_DIR", H_DIR / "reference" / "primary"))
DATABOOK = PRIMARY / "Lenta_Q22026_DATABOOK.xlsx"
IFRS_DIR = PRIMARY / "ifrs_mkpao"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def manifest_sha(name: str) -> str | None:
    """sha256 файла из MANIFEST.md первички (для сверки)."""
    man = PRIMARY / "MANIFEST.md"
    if not man.exists():
        return None
    txt = man.read_text(encoding="utf-8")
    for line in txt.splitlines():
        if name in line:
            m = re.search(r"\b([0-9a-f]{64})\b", line)
            if m:
                return m.group(1)
    return None


def half_key(label: str) -> tuple[str, str] | None:
    """'FY 2019' -> ('FY', '2019'); '1H 2019' / '1Н 2015' (кириллица) -> ('1H', ...)."""
    s = str(label).strip().replace("Н", "H")
    m = re.match(r"(FY|1H)\s*(\d{4})", s)
    return (m.group(1), m.group(2)) if m else None


@lru_cache(maxsize=None)
def databook_rows(sheet: str, max_row: int = 130):
    import openpyxl
    wb = openpyxl.load_workbook(DATABOOK, data_only=True, read_only=True)
    ws = wb[sheet]
    return [tuple(r) for r in ws.iter_rows(min_row=1, max_row=max_row, values_only=True)]


def block_columns(sheet: str, basis_row: int = 7, period_row: int = 8) -> dict:
    """{('IAS 17'|'IFRS 16', 'FY'|'1H', '2019'): col} для листов PL/CF датабука."""
    rows = databook_rows(sheet)
    basis_r, per_r = rows[basis_row - 1], rows[period_row - 1]
    out, basis = {}, None
    for j, (b, p) in enumerate(zip(basis_r, per_r)):
        if b:
            basis = str(b).strip()
        if p is None or basis is None:
            continue
        k = half_key(p)
        if k:
            out[(basis, k[0], k[1])] = j
    return out


def row_by_label(sheet: str, label: str, start: int = 1) -> int:
    """Номер строки (1-based) по подписи в колонке A (начало строки, без регистра)."""
    rows = databook_rows(sheet)
    lab = label.strip().lower()
    for i, r in enumerate(rows, 1):
        if i < start:
            continue
        if r[0] is not None and str(r[0]).strip().lower().startswith(lab):
            return i
    raise KeyError(f"{sheet}: нет строки «{label}»")


def pl_value(label: str, basis: str, kind: str, year: str, sheet: str = "PL") -> float | None:
    """Значение строки PL/CF в млрд руб. (в датабуке — тыс. руб.)."""
    cols = block_columns(sheet)
    key = (basis, kind, year)
    if key not in cols:
        return None
    r = row_by_label(sheet, label)
    v = databook_rows(sheet)[r - 1][cols[key]]
    return None if v is None else float(v) / 1e6


def fq_table() -> dict:
    """Лист 'Financials quarterly' (IAS 17, млн руб.) -> {квартал '2024Q1': {строка: млрд}}."""
    rows = databook_rows("Financials quarterly", 80)
    basis_r, hdr = rows[6], rows[7]
    cols, basis = {}, None
    for j, c in enumerate(hdr):
        if basis_r[j]:
            basis = str(basis_r[j]).strip()
        if c is None or basis != "IAS 17":
            continue
        m = re.match(r"([1-4])Q\s*(\d{4})", str(c).strip())
        if m:
            key = f"{m.group(2)}Q{m.group(1)}"
            if key not in cols:
                cols[key] = j
    want = {
        "revenue": "Total Sales", "hyper": "Hypermarkets", "super": "Supermarkets",
        "conv": "Convenience stores", "utkonos": "Utkonos", "droge": "Drogerie", "remi": "Remi",
        "diy": "Dom Lenta", "other": "Other formats", "wholesale": "Wholesale",
        "gp": "Gross profit", "sga": "Selling, general and administrative expenses",
        "payroll": "Payroll and related taxes", "da": "Depreciationa and Amortization",
        "lease": "Lease Expenses", "ebitdar": "EBITDAR", "ebitda": "EBITDA",
    }
    idx = {}
    for i, r in enumerate(rows[:55], 1):
        if r[0] is None:
            continue
        lab = str(r[0]).strip()
        for k, w in want.items():
            if k not in idx and lab.lower() == w.lower():
                idx[k] = i
    out = {}
    for q, j in cols.items():
        d = {}
        for k, i in idx.items():
            v = rows[i - 1][j]
            d[k] = None if v in (None, "", "-", "–") else float(v) / 1e3
        out[q] = d
    return out


@lru_cache(maxsize=None)
def pdf_text(name: str) -> str:
    """Текст PDF первички (pypdfium2). name — путь относительно PRIMARY."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(PRIMARY / name))
    parts = []
    for i in range(len(pdf)):
        parts.append(f"=== page {i + 1}\n" + pdf[i].get_textpage().get_text_range())
    return "\n".join(parts)


def pdf_number(name: str, pattern: str, group: int = 1) -> float:
    """Число из текста PDF по регулярке; пробелы-разделители тысяч снимаются. тыс. руб. -> млрд."""
    txt = pdf_text(name)
    m = re.search(pattern, txt, flags=re.S)
    if not m:
        raise AssertionError(f"{name}: не найдено {pattern!r}")
    raw = m.group(group).replace("\u00a0", " ").replace(" ", "").replace(",", "")
    return float(raw) / 1e6


def pdf_page_of(name: str, pattern: str) -> int | None:
    txt = pdf_text(name)
    m = re.search(pattern, txt, flags=re.S)
    if not m:
        return None
    return txt[: m.start()].count("=== page ")


def dump(obj, path: Path):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=1, sort_keys=False), encoding="utf-8")


def fmt(x, nd=2):
    return "—" if x is None else f"{x:.{nd}f}".replace(".", ",")

"""Рынок и заголовок: лист «Оценка» книги 1.0 «Ленты» → out/market_out.txt, out/market_out.json.

  A-V11 цена книги: LEGALCLOSEPRICE LENT 18.09.2026 (проверка по трём выгрузкам ISS), конвенция legalclose (D7);
  A-V12 цели инвестдомов: последняя цель каждого дома с датой (research/lent_sellside_targets.csv) + агрегаты;
        аналоги на одной базе (facts/peers.json) + строка LENT на проформенной LTM EBITDA (beta_out.json);
  A-V8  отображение intrinsic: пороги гейта limited_liability (доля опционной части в заголовке) и защиты от скачка;
  копии 850oa: terminal_real_growth, terminal, roll_along_forwards, distress, market_curve_years, ronic_spread, правило полосы.
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import math

import yaml

from common import BOOK_DRAFT, INPUTS, OUT, Log, dump_json, load_iss, load_iss_rows, magnit_dir, norm_cdf, primary_dir, rel, research_dir, sha256_file

L = Log(OUT / "market_out.txt")
R: dict = {"schema": "lenta-book-1.0-valuation-market-v1", "sources": {}}
PRICE_DATE = dt.date(2026, 9, 18)
SHARES_OUT = 115_074_675


def src(name, path):
    R["sources"][name] = {"file": rel(path), "sha256": sha256_file(path)}
    return path


# ============================================================ 1. цена книги
F1 = src("iss_live_2026-09", INPUTS / "iss_LENT_hist_2026-09-01_2026-09-25.json")
F2 = src("iss_research_LENT", research_dir() / "_work05" / "LENT_TQBR.json")
F3 = src("iss_primary_LENT_close", primary_dir() / "moex_iss_lent_corp" / "iss_history_LENT_2026-06-26_2026-09-25.json")
F4 = src("iss_research_MGNT", research_dir() / "_work05" / "MGNT_TQBR.json")
rows1 = {r["TRADEDATE"]: r for r in load_iss_rows(F1)}
rows2 = {r["TRADEDATE"]: r for r in load_iss_rows(F2)}
rows3 = {r["TRADEDATE"]: r for r in load_iss_rows(F3)}
rows4 = {r["TRADEDATE"]: r for r in load_iss_rows(F4)}
k = PRICE_DATE.isoformat()
lc1, lc2 = rows1[k]["LEGALCLOSEPRICE"], rows2[k]["LEGALCLOSEPRICE"]
cl1, cl2, cl3 = rows1[k]["CLOSE"], rows2[k]["CLOSE"], rows3[k]["CLOSE"]
assert lc1 == lc2 == 1619.5 and cl1 == cl2 == cl3 == 1625.5, "цена книги не сходится между выгрузками ISS"
PRICE = lc1
win = ["2026-09-15", "2026-09-16", "2026-09-17", "2026-09-18", "2026-09-21", "2026-09-22", "2026-09-23", "2026-09-24", "2026-09-25"]
R["price"] = {"price": PRICE, "convention": "legalclose", "date": k, "close": cl1, "waprice": rows1[k]["WAPRICE"], "volume": rows1[k]["VOLUME"],
              "window": {d: {"close": rows1[d]["CLOSE"], "legalclose": rows1[d]["LEGALCLOSEPRICE"], "volume": rows1[d]["VOLUME"]} for d in win},
              "mgnt_850oa_check": {d: {"close": rows4[d]["CLOSE"], "legalclose": rows4[d]["LEGALCLOSEPRICE"]} for d in ("2026-09-16", "2026-09-17", "2026-09-18")},
              "mcap_book_bn": PRICE * SHARES_OUT / 1e9}
L("# market_headline.py — рынок и заголовок (лист «Оценка», книга 1.0 «Ленты»)")
L("## 1. цена книги (A-V11, D7)")
L(f"LENT 18.09.2026: LEGALCLOSEPRICE {lc1} (живая выгрузка ISS 28.09 — {rel(F1)}; research/_work05 — {lc2}); CLOSE {cl1} (= primary {rel(F3)}: {cl3}); "
  f"WAPRICE {rows1[k]['WAPRICE']}; объём {rows1[k]['VOLUME']} акций")
for d in win:
    L(f"  {d}: close {rows1[d]['CLOSE']}, legal close {rows1[d]['LEGALCLOSEPRICE']}, объём {rows1[d]['VOLUME']}")
L(f"18.09 — первый день вне IMOEX (последний день в индексе 17.09: legal close 1 573,5, объём ×2–4 к обычному); через неделю legal close 1 746 (+7,8 %)")
L(f"сверка конвенции 850oa: MGNT legal close 16.09 {rows4['2026-09-16']['LEGALCLOSEPRICE']}, 18.09 {rows4['2026-09-18']['LEGALCLOSEPRICE']} — "
  f"1 560 книги 850oa = legal close 16.09, а не 18.09 (конвенция там не воспроизводится; у «Ленты» дата и поле записаны ключами)")
L(f"капитализация по цене книги и акциям в обращении: {PRICE*SHARES_OUT/1e9:.3f} млрд ₽")

# ============================================================ 2. цели инвестдомов
F5 = src("sellside_csv", research_dir() / "lent_sellside_targets.csv")
rows = list(csv.DictReader(open(F5, encoding="utf-8")))
HOUSE_MAP = {"SberCIB": "Сбер", "Сбер": "Сбер", "Т-Инвестиции": "Т-Инвестиции", "Атон": "Атон", "Синара": "Синара", "Цифра брокер": "Цифра брокер",
             "ПСБ": "ПСБ", "АКБФ": "АКБФ", "Эйлер": "Эйлер", "Freedom Global": "Freedom Global", "Совкомбанк": "Совкомбанк",
             "Газпромбанк Инвестиции": "Газпромбанк Инвестиции"}


def house_of(s):
    for key, v in HOUSE_MAP.items():
        if s.startswith(key):
            return v
    return s


latest = {}
excluded = []
for r in rows:
    h = house_of(r["house"])
    if not r["target_rub"]:
        excluded.append((r["date"], r["house"], "без цели"))
        continue
    if "торговая идея" in r["house"]:
        excluded.append((r["date"], r["house"], "торговая идея, закрыта по цели"))
        continue
    if h not in latest or r["date"] > latest[h]["date"]:
        latest[h] = {"house": h, "target": float(r["target_rub"]), "date": r["date"], "rating": r["rating"], "source": r["source"].split(" ; ")[0],
                     "price_on_date": r["price_on_date_legal_close"], "note": r["house"]}
tg = sorted(latest.values(), key=lambda x: x["date"])


def agg(xs, label):
    v = sorted(x["target"] for x in xs)
    n = len(v)
    med = (v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2) if n else None
    return {"label": label, "n": n, "median": med, "mean": sum(v) / n if n else None, "min": v[0] if n else None, "max": v[-1] if n else None,
            "median_vs_book_price": (med / PRICE - 1) if n else None}


AGG = {"all_latest": agg(tg, "последняя цель каждого дома"),
       "after_okey_2026-06-02": agg([x for x in tg if x["date"] >= "2026-06-02"], "после объявления «О'КЕЙ» (02.06.2026)"),
       "after_q2_2026-08-03": agg([x for x in tg if x["date"] >= "2026-08-03"], "после отчёта 2 кв. (03.08.2026)")}
R["sellside"] = {"latest": tg, "excluded": excluded, "agg": AGG, "as_of": "2026-09-28"}
L()
L(f"## 2. цели инвестдомов (A-V12; справочно, в модель не входят; {rel(F5)})")
for x in tg:
    L(f"  {x['date']} {x['house']}: {x['target']:.2f} ₽ ({x['rating']}); цена на дату {x['price_on_date']}; {x['source']}")
for e in excluded:
    L(f"  исключено: {e[0]} {e[1]} — {e[2]}")
for a in AGG.values():
    L(f"{a['label']}: n {a['n']}, медиана {a['median']:.0f}, среднее {a['mean']:.1f}, {a['min']:.2f}–{a['max']:.0f}; медиана к цене книги {a['median_vs_book_price']*100:+.1f} %")
L("пересмотры после «О'КЕЙ» и 2 кв.: вниз — Т-Инвестиции (2 800 → 2 000, «держать»), Эйлер (3 300 → 2 300), АКБФ (1 639 → 1 474); вверх — Цифра, ПСБ, Сбер")

# ============================================================ 3. аналоги на одной базе
FP = BOOK_DRAFT / "facts" / "peers.json"
BO = OUT / "beta_out.json"
peers = json.loads(src("facts_peers", FP).read_text(encoding="utf-8"))
beta_out = json.loads(src("beta_out", BO).read_text(encoding="utf-8"))
lent = peers["lent"]
eb_pf = beta_out["pro_forma_ltm"]["center"]["ebitda"]
eb_pf_q = beta_out["pro_forma_ltm"]["quantiles"]["ebitda"]
ev_l = lent["ev"]
PEERS = [{"name": p["name"], "ev_ebitda": round(p["ev_ebitda"], 3), "basis": "IAS 17", "as_of": "2026-09-18",
          "note": f"EV {p['ev']:.3f} / EBITDA LTM {p['ebitda_ltm']:.3f}; ЧД/EBITDA {p['nd_ebitda']:.3f}; маржа {p['ebitda_margin']*100:.2f} %"} for p in peers["peers"]]
PEERS.append({"name": "Лента (отчётная LTM)", "ev_ebitda": round(ev_l / lent["ebitda_ltm_reported"], 3), "basis": "IAS 17", "as_of": "2026-09-18",
              "note": f"EV {ev_l:.3f} / {lent['ebitda_ltm_reported']:.3f}; завышен: долг «О'КЕЙ» в ЧД целиком, EBITDA — месяц"})
PEERS.append({"name": "Лента (проформа LTM)", "ev_ebitda": round(ev_l / eb_pf, 3), "basis": "IAS 17", "as_of": "2026-09-18",
              "note": f"EV {ev_l:.3f} / EBITDA проформа {eb_pf:.2f} (P10–P90 {eb_pf_q['0.1']:.2f}–{eb_pf_q['0.9']:.2f} → {ev_l/eb_pf_q['0.9']:.2f}–{ev_l/eb_pf_q['0.1']:.2f}×); "
                      f"расчёт листа valuation до канона margin/facts"})
R["peers"] = PEERS
L()
L(f"## 3. аналоги на одной базе IAS 17 (EV = капитализация по legal close 18.09 + ЧД 30.06 + дивиденды после баланса; {rel(FP)})")
for p in PEERS:
    L(f"  {p['name']}: {p['ev_ebitda']:.3f}× — {p['note']}")

# ============================================================ 4. отображение intrinsic: гейт limited_liability и защита от скачка
def call(x, s):
    """Колл Мертона при ставке 0 в долях D: C/D как функция V0/D = x и σ√T = s."""
    d1 = math.log(x) / s + s / 2
    return x * norm_cdf(d1) - norm_cdf(d1 - s)


D_BN = 131.0          # требования на дату книги ≈ ЧД 117,4 + строки моста (пут «Реми» 5,2, «ОБИ» 4,1, LTIP 2,0, НДУ 1,0, прочее) — порядок, лист financing-bridge
b3w = beta_out["own_betas"]["MCFTR|W|3"]
ev_book = beta_out["rho_check"]["E_over_V_book"]
sig_v = b3w["vol_i"] * ev_book
TAB = []
for x in (1.1, 1.2, 1.25, 1.3, 1.4, 1.5, 1.75, 2.0, 2.5):
    row = {"x": x}
    for s in (0.20, 0.25, 0.30, 0.35):
        opt = call(x, s) - (x - 1)
        row[f"s{s}"] = {"opt_share_of_intrinsic": opt / (x - 1), "opt_rub": opt * D_BN * 1e9 / SHARES_OUT}
    TAB.append(row)
LL = {"v0_to_d_min": 1.3, "max_share": 0.10}
JG = {"median_pct": 0.08, "v0_pct": 0.10}
lev = beta_out["ebitdar_route"]["lenta_V"] / beta_out["ebitdar_route"]["lenta_E_book"]
R["headline"] = {"method": "intrinsic", "print_step": 50, "diagnostics": "median", "limited_liability": LL, "jump_guard": JG,
                 "option_table": TAB, "sigma_V_weekly_3y": sig_v, "D_bn_assumed": D_BN, "ev_over_e_book": lev}
L()
L("## 4. отображение V0 → цена: intrinsic (D4); пороги гейта и защиты")
L(f"σ_V ≈ σ_E LENT (недели, 3 г.) {b3w['vol_i']*100:.1f} % × E/V {ev_book:.3f} = {sig_v*100:.1f} %; срок долга 1,5–2 года → σ√T ≈ {sig_v*math.sqrt(1.5):.2f}–{sig_v*math.sqrt(2):.2f}; "
  f"3-мес. σ_E 39 % дала бы до ≈{0.39*ev_book*math.sqrt(2):.2f}")
L("V0/D | опционная часть к внутренней (₽ на акцию при D ≈ 131 млрд) при σ√T 0,20 / 0,25 / 0,30 / 0,35")
for r in TAB:
    L(f"{r['x']:.2f} | " + " / ".join(f"{r[f's{s}']['opt_share_of_intrinsic']*100:.1f} % ({r[f's{s}']['opt_rub']:.0f} ₽)" for s in (0.20, 0.25, 0.30, 0.35)))
L(f"limited_liability: v0_to_d_min {LL['v0_to_d_min']} — ниже опционная часть превышает ≈40 ₽ (≈шаг печати 50 ₽) при σ√T 0,30; max_share {LL['max_share']} — "
  f"P10 полосы (нижний печатаемый квантиль) обязан лежать в зоне V0 ≥ 1,3·D")
L(f"jump_guard: median_pct {JG['median_pct']}, v0_pct {JG['v0_pct']}. EV/E на дату книги {lev:.2f}: порог медианы в долях цены при малом рычаге "
  f"защищает слабо — пороги под рычаг «Ленты»; законный дрейф медианы без новой книги/фактов — перекат по форвардам ≈ r_E/365 ≈ 0,04 %/день (<1 % за две недели); "
  f"8 % — ниже верха D4 (10 %); V0 10 % — ловит перенос статьи между V0 и D при устойчивой медиане")

# ============================================================ 5. копии 850oa (правила)
TPL = src("magnit_template", magnit_dir() / "data" / "assumptions" / "assumptions_template.yaml")
tpl = yaml.safe_load(TPL.read_text(encoding="utf-8").replace("#{{WORLDS}}", ""))
v = tpl["valuation"]
COPY = {k: v[k] for k in ("terminal_real_growth", "roll_along_forwards", "distress", "market_curve_years", "ronic_spread", "terminal")}
COPY["uncertainty"] = {k: v["uncertainty"][k] for k in ("draws", "seed", "quantiles", "median_draws", "reverse_bounds", "median_refine")}
COPY["headline"] = {"print_step": v["headline"]["print_step"], "diagnostics": v["headline"]["diagnostics"]}
R["copied_from_850oa"] = COPY
L()
L(f"## 5. правила, взятые из 850oa без изменений ({rel(TPL)})")
for kk, vv in COPY.items():
    L(f"  {kk}: {vv}")

# ============================================================ 6. A-V4: ROIC «Ленты» (IAS 17) — сверка RONIC − r = 0
import openpyxl  # noqa: E402

DB = src("databook", primary_dir() / "Lenta_Q22026_DATABOOK.xlsx")
wb = openpyxl.load_workbook(DB, data_only=True, read_only=True)
pl = [list(r) for r in wb["PL"].iter_rows(values_only=True)]
bs = [list(r) for r in wb["BS"].iter_rows(values_only=True)]
ph, bh = pl[7], bs[7]
roic = []
for y in range(2018, 2026):
    jp, jb, jb0 = ph.index(f"FY {y}") if f"FY {y}" in ph else None, bh.index(f"FY {y}"), bh.index(f"FY {y-1}")
    if jp is None:
        continue
    ebit = pl[31][jp] / 1e6                     # операционная прибыль до обесценения (стр. 32)
    tax = abs(pl[44][jp]) if isinstance(pl[44][jp], float) else None
    ic = lambda j: (bs[50][j] + bs[53][j] + bs[61][j] - bs[33][j]) / 1e6   # капитал + займы − деньги
    ic_avg = (ic(jb) + ic(jb0)) / 2
    t = 0.20 if y < 2025 else 0.25
    roic.append({"year": y, "ebit_pre_impairment": ebit, "ic_avg": ic_avg, "roic_after_tax": ebit * (1 - t) / ic_avg, "tax_rate_statutory": t})
R["roic_ias17"] = roic
L()
L("## 6. A-V4 (сверка): ROIC «Ленты» IAS 17 = операционная прибыль до обесценения × (1 − ставка) / средний (капитал + займы − деньги)")
for r in roic:
    L(f"  {r['year']}: EBIT {r['ebit_pre_impairment']:.2f}; IC {r['ic_avg']:.1f}; ROIC {r['roic_after_tax']*100:.1f} %")
L("ROIC номинальный 2018–2025 против r_u той же эпохи (ОФЗ 7–15 % + премия ≈3 п.п.): предельная отдача сделок не выше r → RONIC − r = 0 (как 850oa), ось не нужна")


# ============================================================ 7. история EV/EBITDA «Ленты» (IAS 17) — коридор гейта ev_ebitda
lent_px = load_iss(research_dir() / "_work05" / "LENT_TQBR.json")
lnta_px = {kk: vv * 5 for kk, vv in load_iss(research_dir() / "_work05" / "LNTA_TQBR.json").items()}
hist = []
for y in range(2014, 2026):
    ye = dt.date(y, 12, 31)
    series = lnta_px if ye < dt.date(2022, 3, 24) else lent_px
    dd = max(kk for kk in series if kk <= ye)
    sh = 97_585_932 if ye < dt.date(2022, 6, 21) else 115_985_197
    jp, jb = ph.index(f"FY {y}") if f"FY {y}" in ph else None, bh.index(f"FY {y}")
    if jp is None:
        continue
    eb = pl[48][jp] / 1e6
    ndv = (bs[53][jb] + bs[61][jb] - bs[33][jb]) / 1e6
    mc = series[dd] * sh / 1e9
    hist.append({"year": y, "price_date": dd.isoformat(), "price": series[dd], "mcap": mc, "nd": ndv, "ebitda": eb, "ev_ebitda": (mc + ndv) / eb})
R["ev_ebitda_history"] = hist
L()
L("## 7. история EV/EBITDA «Ленты» IAS 17 на конец года (цена ISS: LNTA×5 до 2022, LENT; ЧД и EBITDA — датабук; акции выпущенные) — коридор гейта ev_ebitda")
for r in hist:
    L(f"  {r['year']}: цена {r['price']:.1f} ({r['price_date']}); капитализация {r['mcap']:.1f}; ЧД {r['nd']:.1f}; EBITDA {r['ebitda']:.2f}; EV/EBITDA {r['ev_ebitda']:.2f}×")
e1 = [r["ev_ebitda"] for r in hist if r["year"] <= 2021]
e2 = [r["ev_ebitda"] for r in hist if r["year"] >= 2022]
L(f"2014–2021 (ОФЗ 6–10 %): {min(e1):.1f}–{max(e1):.1f}×; 2022–2025 (ОФЗ 8–16 %): {min(e2):.1f}–{max(e2):.1f}×; 18.09.2026: отчётная 3,60×, проформа 3,48×; "
  f"X5 2,82×, Магнит 3,48× — предложение коридора: 2,5–4,5× для H и M, 3,5–6,0× для N")

L.save()
dump_json(R, OUT / "market_out.json")

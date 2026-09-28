"""Выписка X5 (IAS 17) для операционного рычага и разрычаживания аренды → inputs/x5_ias17.json.

Источник: датабук X5 «2026_08_financial_and_operating_results_q2_2026.xlsx» (публичный файл X5; тот же, что в
листе беты 850oa и в facts/peers.json). Путь — переменная X5_DATABOOK (по умолчанию <handoff>/../magnit-model/
data/x5/books/<файл>). Выписка сохраняется с sha256 книги и адресами ячеек, поэтому beta.py работает и без xlsx.
Берутся кварталы 3 кв. 2025 – 2 кв. 2026 (LTM на 30.06.2026) и все даты обязательств по аренде.
"""
from __future__ import annotations

import datetime as dt
import os
from pathlib import Path

import openpyxl

from common import HANDOFF, INPUTS, dump_json, sha256_file

FNAME = "2026_08_financial_and_operating_results_q2_2026.xlsx"
QUARTERS = ["3 КВ. 2025", "4 КВ. 2025", "1 КВ. 2026", "2 КВ. 2026"]


def book_path() -> Path:
    return Path(os.environ.get("X5_DATABOOK", HANDOFF.parent / "magnit-model" / "data" / "x5" / "books" / FNAME))


def ias17_quarter_cols(ws, hdr_row=5, basis_row=4):
    """Колонки кварталов блока IAS 17 (первое вхождение подписи квартала с базисом IAS 17)."""
    out = {}
    hdr = [c.value for c in ws[hdr_row]]
    bas = [c.value for c in ws[basis_row]]
    for j, (h, b) in enumerate(zip(hdr, bas)):
        if isinstance(h, str) and "КВ." in h and b == "IAS 17" and h not in out:
            out[h] = j
    return out


def find_row(ws, label, col=1, start=1, end=80):
    for i in range(start, end + 1):
        v = ws.cell(row=i, column=col + 1).value
        if isinstance(v, str) and v.strip() == label:
            return i
    raise KeyError(label)


def main():
    p = book_path()
    wb = openpyxl.load_workbook(p, data_only=True)
    res = {"source": FNAME, "sha256": sha256_file(p), "unit": "млрд руб.", "basis": "IAS 17", "quarters": {}, "cells": {}}
    # P&L: выручка, валовая прибыль
    pl = wb["Profit and Loss"]
    q_pl = ias17_quarter_cols(pl)
    r_rev, r_gp = find_row(pl, "Выручка"), find_row(pl, "Валовая прибыль")
    # SG&A: аренда (IAS 17)
    sga = wb["SG&A"]
    q_sga = ias17_quarter_cols(sga)
    r_rent = find_row(sga, "Расходы на аренду")
    # EBITDA (IAS 17)
    eb = wb["EBITDA"]
    q_eb = ias17_quarter_cols(eb)
    r_eb = find_row(eb, "EBITDA", start=6)
    for q in QUARTERS:
        res["quarters"][q] = {
            "revenue": pl.cell(row=r_rev, column=q_pl[q] + 1).value / 1000,
            "gross_profit": pl.cell(row=r_gp, column=q_pl[q] + 1).value / 1000,
            "rent": sga.cell(row=r_rent, column=q_sga[q] + 1).value / 1000,
            "ebitda": eb.cell(row=r_eb, column=q_eb[q] + 1).value / 1000,
        }
        res["cells"][q] = {"revenue": f"Profit and Loss!r{r_rev}c{q_pl[q]+1}", "gross_profit": f"Profit and Loss!r{r_gp}c{q_pl[q]+1}",
                           "rent": f"SG&A!r{r_rent}c{q_sga[q]+1}", "ebitda": f"EBITDA!r{r_eb}c{q_eb[q]+1}"}
    # Debt: обязательства по аренде и ЧД до МСФО 16 по датам
    dbt = wb["Debt"]
    dates = [c.value for c in dbt[5]]
    r_l = find_row(dbt, "Обязательства по аренде")
    r_nd = find_row(dbt, "Чистый долг до применения МСФО (IFRS) 16")
    lease, nd = {}, {}
    for j, d in enumerate(dates):
        if isinstance(d, dt.datetime):
            k = d.date().isoformat()
            lease[k] = dbt.cell(row=r_l, column=j + 1).value / 1000
            nd[k] = dbt.cell(row=r_nd, column=j + 1).value / 1000
    res["lease_liabilities"] = dict(sorted(lease.items()))
    res["net_debt_ias17"] = dict(sorted(nd.items()))
    res["cells"]["lease_liabilities"] = f"Debt!r{r_l}"
    res["cells"]["net_debt_ias17"] = f"Debt!r{r_nd}"
    ltm = {k: sum(res["quarters"][q][k] for q in QUARTERS) for k in ("revenue", "gross_profit", "rent", "ebitda")}
    res["ltm_2026H1"] = ltm
    dump_json(res, INPUTS / "x5_ias17.json")
    print("X5 LTM 30.06.2026:", {k: round(v, 3) for k, v in ltm.items()})
    print("ВП/EBITDA", round(ltm["gross_profit"] / ltm["ebitda"], 4), " ВП/EBITDAR", round(ltm["gross_profit"] / (ltm["ebitda"] + ltm["rent"]), 4))
    print("sha256", res["sha256"])


if __name__ == "__main__":
    main()

"""Сборщики источников слоя индикаторов Ленты (D15) и дисциплина их реестра.

Состав — минимальный честный пакет D15:

* рынок: `moex_quote` (котировка эмитента — КРИТИЧЕСКИЙ источник; аналоги X5 и
  MGNT — справочно, на одной базе), `moex_curve` (бескупонная кривая и ОФЗ-ИН —
  критический), `moex_security` (уровень листинга и объём выпуска акции, раз в
  сутки: смена — плашка и гейт «пересмотреть g»), `cbr_key_rate`, `bonds`
  (облигации эмитентов группы по кодам ISS — справочно);
* эмитент: `lenta_disclosure` (лента существенных фактов на lentagroup.ru,
  шаблон Zebra: JSON `App = {...}` внутри страницы — как magnit.com у 850oa),
  `lenta_databook` (датабук со страницы публикаций: защищённый источник,
  каждая новая версия хранится с sha256, сигнал «вышел отчёт», квартальный
  P&L IAS 17 — в ряды с винтажами), `trudvsem_vacancies` (вакансии групп
  юрлиц на «Работе России» — справочный ряд вилок; ФИО и контакты вырезаются
  ДО записи на диск — задокументированное исключение из «сырое до разбора»).

Реестр сборщиков (`COLLECTORS`) и его дисциплина — внизу модуля: критические
источники (`CRITICAL`), невосполнимые (`IRRECOVERABLE`), раз в сутки
(`ONCE_A_DAY`), вправе молчать (`MAY_PARSE_NOTHING`), расписание (`DAILY`,
`WEEKLY`). Имя эмитента, тикер, адреса и ИНН — `indicators.issuer`.

**Вежливость** (VPS общий для всех панелей владельца; бан одного IP ослепит
все): пауза между обращениями к одному хосту — `indicators.http`
(`DEFAULT_THROTTLE`, 2 с), без параллельности. Бюджет запросов в сутки при
двух тактах (утренний `collect` и вечерний `daily`): ISS ≈ 14, lentagroup.ru
≤ 10 (две страницы на такт и ≤ 3 тела фактов на такт; датабук — только новая
версия или раз в неделю), «Работа России» ≤ 12 страниц раз в сутки.

Что собирать нельзя и почему (проверено у 850oa и в пробах Ленты, повторно не
предлагать — study/04 §6, research/06 §8): hh.ru (API закрыт без токена),
e-disclosure (капча, бан IP), lenta.com, rabota.lenta.com и monetka.ru (Qrator —
защиту от ботов не обходим), Чек Индекс и СберИндекс (не улучшают прогноз),
цены с полки (robots.txt), `?years=` на lentagroup.ru (robots.txt), cron на
cmegroup.com с VPS (бан IP утащит все панели владельца).
"""

from __future__ import annotations

import hashlib
import html
import inspect
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Callable

from indicators import issuer
from indicators.http import FetchError, RawSink, Response, fetch
from indicators.store import Point, Store

UTC_TODAY = lambda: datetime.now(timezone.utc).date().isoformat()

# Метки времени Zebra (lentagroup.ru) — секунды UTC; день публикации — по Москве.
MSK = timezone(timedelta(hours=3))


@dataclass(frozen=True)
class Collected:
    source: str
    channel: int
    series: dict[str, list[Point]]
    raw: list[tuple[str, bytes, Response]]
    """(имя файла, байты, ответ) — всё сохраняется до разбора."""
    note: str = ""
    error: str = ""
    """Отказ, случившийся ПОСЛЕ того, как часть дня уже собрана.

    Исключение здесь не годится: `run_collectors` ловит его до сохранения, и
    вместе с отказом на пол пропадает всё, что сборщик успел разобрать (у
    850oa обход вакансий терял ряд невосстановимого дня из-за одной локации,
    отдавшей мусор). Поэтому сборщик возвращает собранное и НАЗЫВАЕТ отказ, а
    статус источника всё равно становится «ОШИБКА» — тревога уходит кодом
    возврата юнита, как и прежде.
    """
    idle: str = ""
    """Источник САМ сказал, что наблюдений сегодня нет, — и это не сбой.

    Пустой ответ от мусора в ответе отличает только сборщик (биржа до начала
    торгов: последняя сделка пуста, сделок 0). Причина называется здесь, и
    `run_collectors` пишет её в статус «ок» без точек; без неё «разобрано 0
    рядов» остаётся отказом, как прежде.
    """


# ------------------------------------------------- канал 1: ставки и долг


KEY_RATE_FIRST_DAY = date(2013, 9, 13)
"""Первый день существования ключевой ставки. Глубокая догрузка ходит сюда."""

KEY_RATE_DEPTH_NEEDED = date(2024, 12, 31)
"""Насколько глубоко ряд обязан доходить, чтобы бэктест канала процентов
имел смысл: первое проверяемое полугодие — 1П2025."""


def collect_cbr_key_rate(*, since: date | None = None, store: Store | None = None) -> Collected:
    """Ключевая ставка ЦБ. SOAP отдаёт XML вместо вёрстки — парсер не ломается
    при редизайне сайта.

    Суточный прогон просит год: ответ маленький, а архив не растёт. Глубокая
    история нужна для бэктеста канала процентов: без неё средняя ключевая за
    1П2025 считалась по первому наблюдению ряда и выходила 18,00 % вместо
    20,88 %. Ставка не пересматривается задним числом, поэтому догрузка не
    нарушает точечности винтажей.

    Догрузка САМОЛЕЧАЩАЯСЯ: если переданный архив ещё не доходит до нужной
    глубины, запрашивается вся история. Разовая команда для этого тоже есть,
    но полагаться на то, что её кто-то вспомнит, нельзя — на свежей машине
    ряд молча оказывался коротким, и тесты, зависящие от него, роняли боевой
    конвейер.
    """
    today = date.today()
    if since is None and store is not None:
        series = store.load("cbr.key_rate")
        oldest = min((p.period for p in series.points), default=None) if series else None
        if oldest is None or oldest[:10] > KEY_RATE_DEPTH_NEEDED.isoformat():
            since = KEY_RATE_FIRST_DAY
    # Сырое — на диск СРАЗУ по получении (S1.4), а не после разбора: разбор
    # чинится задним числом, скачанные байты — нет.
    sink = RawSink(store, "cbr") if store is not None else None
    start = since or date(today.year - 1, today.month, 1)
    envelope = f"""<?xml version="1.0" encoding="utf-8"?>
<soap:Envelope xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/">
 <soap:Body><KeyRate xmlns="http://web.cbr.ru/">
  <fromDate>{start.isoformat()}T00:00:00</fromDate>
  <ToDate>{today.isoformat()}T00:00:00</ToDate>
 </KeyRate></soap:Body></soap:Envelope>"""
    response = fetch(
        "https://www.cbr.ru/DailyInfoWebServ/DailyInfo.asmx",
        data=envelope.encode("utf-8"),
        headers={"Content-Type": "text/xml; charset=utf-8",
                 "SOAPAction": "http://web.cbr.ru/KeyRate"},
        sink=sink, name="key_rate.xml",
    )
    points = []
    for dt, rate in re.findall(r"<DT>([^<]+)</DT>\s*<Rate>([^<]+)</Rate>", response.text):
        points.append(Point(period=dt[:10], value=float(rate) / 100.0,
                            fetched_at=response.fetched_at, source_sha256=response.sha256))
    return Collected("cbr", 1, {"cbr.key_rate": points},
                     [] if sink is not None else [("key_rate.xml", response.body, response)])


def collect_moex_curve(*, store=None) -> Collected:
    """Кривая бескупонной доходности ОФЗ (ZCYC) и реальные доходности ОФЗ-ИН.

    Пара «номинальная кривая / линкеры» даёт вменённую инфляцию — ту самую,
    на которой стоит мир M. Считается точно, `(1+n)/(1+r)−1`, а не разностью.

    Обе страницы ложатся на диск СРАЗУ по получении (S1.4). Прежде HTML вместо
    JSON у второй страницы уносил с собой и первую, УДАЧНУЮ: сырое сохранялось
    вызывающим после возврата сборщика, а до возврата дело не доходило.
    """
    sink = RawSink(store, "moex") if store is not None else None
    raw, series = [], {}

    zcyc = fetch("https://iss.moex.com/iss/engines/stock/zcyc.json?iss.meta=off",
                 sink=sink, name="zcyc.json")
    if sink is None:
        raw.append(("zcyc.json", zcyc.body, zcyc))
    block = zcyc.json()["yearyields"]
    columns = {c: i for i, c in enumerate(block["columns"])}
    for row in block["data"]:
        tenor = row[columns["period"]]
        series.setdefault(f"moex.zcyc.{tenor:g}y", []).append(Point(
            period=row[columns["tradedate"]], value=row[columns["value"]] / 100.0,
            fetched_at=zcyc.fetched_at, source_sha256=zcyc.sha256))

    url = (f"https://iss.moex.com/iss/engines/stock/markets/bonds/boards/TQOB/securities.json"
           f"?iss.meta=off&securities={','.join(OFZ_IN_SECIDS)}"
           f"&iss.only=marketdata,marketdata_yields&marketdata.columns=SECID,YIELD"
           f"&marketdata_yields.columns=SECID,TRADEMOMENT")
    ofz_in = fetch(url, sink=sink, name="ofz_in.json")
    if sink is None:
        raw.append(("ofz_in.json", ofz_in.body, ofz_in))
    linkers, skipped = ofz_in_series(ofz_in.json(), fetched_at=ofz_in.fetched_at,
                                     sha256=ofz_in.sha256)
    series.update(linkers)
    return Collected("moex", 1, series, raw,
                     note="ОФЗ-ИН без точки: " + "; ".join(skipped) if skipped else "")


# Выпуски ОФЗ-ИН рецепта миров (`data/assumptions/worlds_inputs.yaml`,
# `ofz_in.bonds`: 52002–52005). Выпуск, добавленный во входы без строки здесь,
# ловит тест (`tests/test_ofz_in.py`): рецепту 1.4 нужны ВСЕ выпуски входов.
OFZ_IN_SECIDS = ("SU52002RMFS1", "SU52003RMFS9", "SU52004RMFS7", "SU52005RMFS4")


def ofz_in_series(payload: dict, *, fetched_at: str, sha256: str) -> tuple[dict, list[str]]:
    """Реальные доходности ОФЗ-ИН, датированные ДНЁМ ТОРГОВ, и пропущенные выпуски.

    `marketdata.YIELD` — доходность последней сделки (та же величина, что в
    записи книги 1.4: LAST 18.09.2026), и относится она ко дню этой сделки,
    а не ко дню съёма. Прежде точка датировалась `UTC_TODAY()`: съём в
    воскресенье 20.09 писал «2026-09-20» пятничный LAST 18.09, а ранний такт
    сбора (06:10 МСК, до сессии) — вчерашнюю сделку сегодняшним числом.
    Рецепту миров 1.4 дата доходностей обязана совпасть с датой кривой
    (`worlds_recipe.py`, защита по дате), поэтому дата берётся из ответа ISS:
    `marketdata_yields.TRADEMOMENT` — момент последней сделки, «ГГГГ-ММ-ДД
    ЧЧ:ММ:СС» по Москве, как `tradedate` у кривой. Без момента сделки (или с
    нулевой доходностью — сделки не было) точка не пишется: дата съёма на её
    месте и была ошибкой.
    """
    md, my = payload["marketdata"], payload.get("marketdata_yields") or {}
    idx = {c: i for i, c in enumerate(md["columns"])}
    moments = {}
    if my.get("columns"):
        iy = {c: i for i, c in enumerate(my["columns"])}
        moments = {row[iy["SECID"]]: row[iy["TRADEMOMENT"]] for row in my["data"]}
    series, skipped = {}, []
    for row in md["data"]:
        secid, value = row[idx["SECID"]], row[idx["YIELD"]]
        moment = str(moments.get(secid) or "")
        try:
            day = date.fromisoformat(moment[:10]).isoformat()
        except ValueError:
            day = ""
        if not value or not day:
            skipped.append(f"{secid}: " + ("нет доходности сделки" if not value
                                           else f"нет даты сделки ({moment or 'TRADEMOMENT пуст'})"))
            continue
        series[f"moex.ofz_in.{secid}"] = [Point(
            period=day, value=value / 100.0, fetched_at=fetched_at, source_sha256=sha256,
            note=f"сделка {moment} МСК")]
    return series, skipped


def collect_moex_quote(*, store=None) -> Collected:
    """Котировка эмитента (`issuer.TICKER`) и аналогов (`issuer.PEER_TICKERS`).
    Бесплатный ISS отдаёт с задержкой 15 минут — опрашивать чаще раза в пять
    минут бессмысленно и невежливо.

    Страница ложится на диск СРАЗУ по получении (S1.4): HTML вместо JSON не
    должен стоить дня наблюдений по КРИТИЧЕСКОМУ источнику.

    Биржа ещё не торгует — не сбой (тревога 03.10.2026, суббота): в 07:00 МСК
    ISS обнуляет таблицу, и до первой сделки дня `LAST` пуст у всех бумаг при
    `NUMTRADES` = 0. В будни утренний такт приходит уже на утреннюю сессию, в
    выходной и в биржевой праздник — на пустую таблицу. Такой ответ — «новых
    точек нет» (`Collected.idle`), выпуск считается на последней принятой цене.
    Пустой `LAST` при ненулевом числе сделок, пропавшая строка эмитента или
    ответ без `NUMTRADES` остаются отказом.
    """
    sink = RawSink(store, "moex") if store is not None else None
    tickers = tuple(dict.fromkeys((issuer.TICKER, *issuer.PEER_TICKERS)))
    url = (f"https://iss.moex.com/iss/engines/stock/markets/shares/boards/TQBR/securities.json"
           f"?iss.meta=off&securities={','.join(tickers)}&iss.only=marketdata"
           f"&marketdata.columns=SECID,LAST,NUMTRADES,TRADINGSTATUS,UPDATETIME")
    response = fetch(url, sink=sink, name="quotes.json")
    md = response.json()["marketdata"]
    idx = {c: i for i, c in enumerate(md["columns"])}
    today = UTC_TODAY()
    series = {}
    for row in md["data"]:
        if row[idx["LAST"]] is not None:
            series[f"moex.price.{row[idx['SECID']]}"] = [Point(
                period=today, value=float(row[idx["LAST"]]),
                fetched_at=response.fetched_at, source_sha256=response.sha256)]
    idle = ""
    if not series and "NUMTRADES" in idx:
        rows = {row[idx["SECID"]]: row for row in md["data"]}
        if issuer.TICKER in rows and all(row[idx["NUMTRADES"]] == 0 for row in rows.values()):
            idle = (f"торгов сегодня ещё не было (последняя сделка пуста, сделок 0 у "
                    f"{len(rows)} бумаг) — новых точек нет")
    return Collected("moex", 1, series,
                     [] if sink is not None else [("quotes.json", response.body, response)],
                     idle=idle)


# ------------------------------------------ описание бумаги: листинг и выпуск


# Поля карточки бумаги ISS, которые собираются: уровень листинга (с 12.08.2026
# акция в третьем уровне, research/01) и объём выпуска в штуках (84 млн
# объявленных акций — риск размещения, D5 (г)). Смена любого — плашка на
# витрине и гейт «пересмотреть g» (P4 читает `security_changes`).
SECURITY_FIELDS = {"LISTLEVEL": "listlevel", "ISSUESIZE": "issuesize"}


def security_series_id(field: str) -> str:
    return f"moex.security.{issuer.TICKER}.{SECURITY_FIELDS[field]}"


def security_series(payload: dict, *, day: str, fetched_at: str, sha256: str) -> dict:
    """Карточка бумаги (`description`) → ряды уровня листинга и объёма выпуска."""
    block = payload["description"]
    columns = {c: i for i, c in enumerate(block["columns"])}
    values = {row[columns["name"]]: row[columns["value"]] for row in block["data"]}
    series = {}
    for field in SECURITY_FIELDS:
        raw = values.get(field)
        try:
            value = float(raw)
        except (TypeError, ValueError):
            continue
        series[security_series_id(field)] = [Point(
            period=day, value=value, fetched_at=fetched_at, source_sha256=sha256,
            note=f"{field} = {raw}")]
    return series


def collect_moex_security(*, store=None) -> Collected:
    """Уровень листинга и объём выпуска акции эмитента — раз в сутки (`ONCE_A_DAY`).

    Карточка меняется раз в годы; второй запрос за день не приносит ничего,
    кроме лишнего обращения к ISS.
    """
    sink = RawSink(store, "moex") if store is not None else None
    url = (f"https://iss.moex.com/iss/securities/{issuer.TICKER}.json"
           "?iss.meta=off&iss.only=description")
    response = fetch(url, sink=sink, name=f"security_{issuer.TICKER}.json")
    series = security_series(response.json(), day=UTC_TODAY(),
                             fetched_at=response.fetched_at, sha256=response.sha256)
    return Collected("moex", 1, series,
                     [] if sink is not None else
                     [(f"security_{issuer.TICKER}.json", response.body, response)])


def security_changes(store) -> list[dict]:
    """Смены карточки бумаги: последнее значение против предыдущего ДРУГОГО.

    Пустой список — карточка не менялась (или ряда ещё нет). Для плашки на
    витрине и гейта «пересмотреть g» (D15): смена уровня листинга или объёма
    выпуска — повод пересмотреть дисконт за управление, а не число модели.
    """
    out = []
    for field in SECURITY_FIELDS:
        series = store.load(security_series_id(field)) if store is not None else None
        points = sorted((p for p in (series.points if series else []) if p.status == "ok"
                         and p.value is not None), key=lambda p: (p.period, p.fetched_at))
        if not points:
            continue
        last = points[-1]
        before = next((p for p in reversed(points[:-1]) if p.value != last.value), None)
        if before is not None:
            out.append(dict(series=series.id, field=field, before=before.value,
                            after=last.value, since=_first_day_of_value(points, last.value),
                            previous_day=before.period))
    return out


def _first_day_of_value(points: list[Point], value: float) -> str:
    """Первый день последнего непрерывного отрезка с этим значением."""
    day = points[-1].period
    for point in reversed(points):
        if point.value != value:
            break
        day = point.period
    return day


# ----------------------------------------- облигации эмитентов группы (ISS)


def bond_issues(payload: dict, emitter_ids) -> dict[str, dict]:
    """Выпуски эмитентов из ответа поиска ISS: {SECID: сведения} — только в обращении.

    Фильтр — по КОДУ эмитента ISS, а не по краткому имени: префикс «Лента»
    совпадает с чужими бумагами, а облигации в обращении у группы — выпуски
    ООО «О'КЕЙ» (эмитент 4867); у ООО «Лента» (5997) все выпуски погашены
    (`is_traded = 0`, research/06 §5).
    """
    block = payload["securities"]
    columns = {c: i for i, c in enumerate(block["columns"])}
    wanted = {int(i) for i in emitter_ids}
    out = {}
    for row in block["data"]:
        emitter = row[columns["emitent_id"]]
        if emitter is None or int(emitter) not in wanted:
            continue
        if not row[columns["is_traded"]] or row[columns["primary_boardid"]] != "TQCB":
            continue
        out[row[columns["secid"]]] = dict(shortname=row[columns["shortname"]],
                                          emitter_id=int(emitter))
    return out


def bond_quote_series(payload: dict, issues: dict[str, dict], *, day: str,
                      fetched_at: str, sha256: str) -> dict:
    """Доска TQCB по выпускам группы → цены (доля номинала) и доходности."""
    sec = payload["securities"]
    si = {c: i for i, c in enumerate(sec["columns"])}
    md = payload["marketdata"]
    mi = {c: i for i, c in enumerate(md["columns"])}
    quotes = {row[mi["SECID"]]: row for row in md["data"]}
    series = {}
    for row in sec["data"]:
        secid = row[si["SECID"]]
        if secid not in issues:
            continue
        quote = quotes.get(secid)
        if not quote:
            continue
        offer = row[si["OFFERDATE"]] if "OFFERDATE" in si else None
        note = (f"{row[si['SHORTNAME']]}, погашение {row[si['MATDATE']]}"
                + (f", оферта {offer}" if offer else ""))
        last, yld = quote[mi["LAST"]], quote[mi["YIELD"]] if "YIELD" in mi else None
        if last is not None:
            series[issuer.series(f"bond.{secid}.price")] = [Point(
                period=day, value=float(last) / 100.0, fetched_at=fetched_at,
                source_sha256=sha256, note=note)]
        if yld:
            series[issuer.series(f"bond.{secid}.yield")] = [Point(
                period=day, value=float(yld) / 100.0, fetched_at=fetched_at,
                source_sha256=sha256, note=note)]
    return series


def collect_bonds(*, store=None) -> Collected:
    """Облигации эмитентов группы на TQCB — справочное наблюдение стоимости долга.

    Два шага: поиск выпусков по ИНН каждого эмитента (`issuer.BOND_ISSUERS`,
    ответ фильтруется по коду эмитента), затем одна доска TQCB по найденным
    выпускам. Цена — доля номинала; у флоатеров `YIELD` непригодна как
    кредитный спред — сохраняется всё, используется осознанно. Ряд
    `<префикс>.bonds.traded` (сколько выпусков в обращении) пишется всегда:
    «ни одного выпуска» — тоже наблюдение, а не пустой ответ.

    Набор данных без кодов эмитентов (`bond_issuers` пуст) и с префиксом имени
    (`bond_shortname_prefixes`, сверка порта на книге 850oa) — прежний путь:
    вся доска TQCB и фильтр по префиксу.

    Каждая страница ложится на диск СРАЗУ по получении (S1.4).
    """
    sink = RawSink(store, "moex") if store is not None else None
    raw: list = []
    today = UTC_TODAY()
    if not issuer.BOND_ISSUERS:
        return _collect_bonds_by_prefix(sink, raw, today)

    issues: dict[str, dict] = {}
    last: Response | None = None
    for item in issuer.BOND_ISSUERS:
        url = ("https://iss.moex.com/iss/securities.json?iss.meta=off"
               f"&q={item['inn']}&limit=100&securities.columns=secid,shortname,isin,"
               "emitent_id,emitent_inn,is_traded,primary_boardid")
        name = f"bonds_search_{item['inn']}.json"
        last = fetch(url, sink=sink, name=name)
        if sink is None:
            raw.append((name, last.body, last))
        issues.update(bond_issues(last.json(), [item["emitter_id"]]))

    series: dict = {}
    if issues:
        url = ("https://iss.moex.com/iss/engines/stock/markets/bonds/boards/TQCB/securities.json"
               f"?iss.meta=off&securities={','.join(sorted(issues))}"
               "&iss.only=securities,marketdata"
               "&securities.columns=SECID,SHORTNAME,ISIN,MATDATE,COUPONPERCENT,OFFERDATE,"
               "FACEVALUE,ISSUESIZEPLACED&marketdata.columns=SECID,LAST,YIELD,UPDATETIME")
        last = fetch(url, sink=sink, name="bonds_tqcb.json")
        if sink is None:
            raw.append(("bonds_tqcb.json", last.body, last))
        series.update(bond_quote_series(last.json(), issues, day=today,
                                        fetched_at=last.fetched_at, sha256=last.sha256))
    series[issuer.series("bonds.traded")] = [Point(
        period=today, value=float(len(issues)), fetched_at=last.fetched_at,
        source_sha256=last.sha256,
        note=(", ".join(f"{v['shortname']} ({k})" for k, v in sorted(issues.items()))
              or "выпусков в обращении нет")[:200])]
    return Collected("moex", 1, series, raw,
                     note=f"выпусков в обращении {len(issues)}")


def _collect_bonds_by_prefix(sink, raw: list, today: str) -> Collected:
    """Прежний путь 850oa: вся доска TQCB, выпуск — по префиксу краткого имени."""
    url = ("https://iss.moex.com/iss/engines/stock/markets/bonds/boards/TQCB/securities.json"
           "?iss.meta=off&iss.only=securities,marketdata"
           "&securities.columns=SECID,SHORTNAME,MATDATE,COUPONPERCENT"
           "&marketdata.columns=SECID,LAST,YIELD")
    response = fetch(url, sink=sink, name="bonds.json")
    if sink is None:
        raw.append(("bonds.json", response.body, response))
    data = response.json()
    sec = data["securities"]
    si = {c: i for i, c in enumerate(sec["columns"])}
    md = data["marketdata"]
    mi = {c: i for i, c in enumerate(md["columns"])}
    prices = {r[mi["SECID"]]: r for r in md["data"]}
    series = {}
    for row in sec["data"]:
        name = str(row[si["SHORTNAME"]] or "")
        if not issuer.BOND_SHORTNAME_PREFIXES or not name.startswith(
                issuer.BOND_SHORTNAME_PREFIXES):
            continue
        quote = prices.get(row[si["SECID"]]) or []
        last = quote[mi["LAST"]] if quote else None
        if last is not None:
            series[issuer.series(f"bond.{row[si['SECID']]}.price")] = [Point(
                period=today, value=float(last) / 100.0,
                fetched_at=response.fetched_at, source_sha256=response.sha256,
                note=f"{name}, погашение {row[si['MATDATE']]}")]
    return Collected("moex", 1, series, raw)


# ------------------------------------------------ Zebra: JSON `App = {...}`


class AppBlockError(RuntimeError):
    """Страница Zebra, из которой блок данных не вынимается.

    «Событий нет» и «мы перестали понимать страницу» — разные вещи: смена
    вёрстки убила бы канал беззвучно. Непонятая страница — отказ сборщика,
    а сама страница при этом ложится на диск ДО отказа.
    """


def app_json(page: str) -> dict:
    """Встроенный в страницу JSON `App = {...}` (шаблон Zebra: lentagroup.ru, magnit.com).

    Разбор декодером JSON с позиции маркера, а не регулярным выражением: в
    названиях и текстах полно и скобок, и кавычек.
    """
    marker = "App = "
    if marker not in page:
        raise AppBlockError(f"в теле страницы {len(page)} знаков, но маркера {marker!r} нет — "
                            "вёрстка сменилась")
    start = page.index(marker) + len(marker)
    try:
        value, _ = json.JSONDecoder().raw_decode(page[start:])
    except json.JSONDecodeError as exc:
        raise AppBlockError(f"блок {marker!r} найден, но не разбирается как JSON: {exc}") from exc
    if not isinstance(value, dict):
        raise AppBlockError(f"блок {marker!r} — не объект JSON")
    return value


def _component(app: dict, path: tuple[str, ...]):
    node = app.get("components") or {}
    for key in path:
        if not isinstance(node, dict) or key not in node:
            raise AppBlockError("в App.components нет пути " + ".".join(path))
        node = node[key]
    return node


def msk_day(stamp) -> str | None:
    """Метка Zebra (секунды UTC) → день по Москве, ISO. Не число — None."""
    if not isinstance(stamp, (int, float)) or isinstance(stamp, bool):
        return None
    return datetime.fromtimestamp(stamp, tz=MSK).date().isoformat()


# ------------------------------------------- лента существенных фактов


DISCLOSURE_SOURCE = "lenta_disclosure"

# Классы событий по ОФИЦИАЛЬНОМУ названию сообщения (формулировки положения ЦБ
# меняются редко, в отличие от вёрстки). Порядок важен: побеждает первое
# совпадение («Приобретение … размещённых ценных бумаг» — оферта, а не сделка).
# Имена классов 850oa сохранены (витрина знает их подписи), добавлены классы
# Ленты: результаты, сделки M&A, собственные акции, дивиденды.
DISCLOSURE_KINDS = {
    "coupon": ("Выплаченные доходы", "Начисленные доходы"),
    "redemption": ("Погашение облигаций",),
    "offer": ("Приобретение (наступление оснований для приобретения) размещённых",
              "Приобретение (наступление оснований для приобретения) размещенных"),
    "placement": ("Дата начала размещения", "Начало размещения", "Завершение размещения",
                  "Регистрация выпуска", "регистрация выпуска", "уведомления об итогах выпуска",
                  "о размещении ценных бумаг"),
    "programme": ("программы облигаций", "программу облигаций"),
    "rating": ("рейтинга",),
    "results": ("финансовой отчетности", "финансовой отчётности", "Операционные результаты",
                "финансовых результатов", "ежеквартального отчета", "годового отчета"),
    "dividends": ("дивиденд",),
    "own_shares": ("собственных акций", "подконтрольной эмитенту организацией голосующих акций",
                   "подконтрольной эмитенту организацией договора"),
    "listing": ("допущенных к торгам", "котировальн", "листинг", "делистинг"),
    "record_date": ("на которую определяются лица", "Дата определения (фиксации) лиц"),
    "board": ("совета директоров", "общего собрания", "общих собраний",
              "единственного акционера", "руководящем составе", "о назначении"),
    "ma": ("существенной сделки", "подконтрольной организации",
           "права распоряжаться определенным количеством голосов", "приобретению",
           "приобретает"),
    "material": ("существенное влияние",),
}

# Уточнение класса по ТЕЛУ — для сообщений «о существенном влиянии», за
# которыми на деле стоят пресс-релизы о результатах, сделках и листинге
# (21 из 300 фактов МКПАО за 2024–2026). Сначала — заголовок пресс-релиза
# (первые `HEADLINE_CHARS` знаков содержания): в сообщении о покупке «Молнии»
# есть и «выручка» приобретаемой сети, но заголовок — «ПРИОБРЕТАЕТ». Затем —
# весь текст; там «приобрет» встречается и в пресс-релизе о результатах,
# поэтому сделка проверяется последней.
BODY_NEEDLES = {
    "results": ("выручк", "рентабельност", "результат"),
    "listing": ("уровень листинга", "уровня листинга", "листинг", "котировальн",
                "списка ценных бумаг"),
    "dividends": ("дивиденд",),
    "own_shares": ("собственных акций", "обратный выкуп", "выкупа акций"),
    "ma": ("приобретает", "приобретени", "сделк", "покупк"),
}
HEADLINE_ORDER = ("listing", "dividends", "own_shares", "ma", "results")
BODY_ORDER = ("results", "listing", "dividends", "own_shares", "ma")
HEADLINE_CHARS = 400
REFINED_KINDS = frozenset({"material", "other"})

# Классы для строки «СОБЫТИЕ» в журнале такта (и, по решению P4, для плашки
# выпуска): то, что двигает оценку или требует пересмотра книги.
DISCLOSURE_ALARM_KINDS = ("ma", "own_shares", "dividends", "listing", "offer", "placement")

# Сколько тел сообщений читать за такт. Бюджет lentagroup.ru — ≤ 6 тел в
# сутки (research/06 §7) при двух тактах. Новые события идут первыми;
# история догоняется по три тела за такт и только в окне `BODY_BACKFILL_DAYS`.
MAX_EVENT_BODIES = 3
BODY_BACKFILL_DAYS = 400

EVENT_ID_RE = re.compile(r"/(\d+)/?$")
EVENT_FILE = re.compile(r"event_(\d+)\.json(?:\.gz)?$")
FEED_FILE = re.compile(r"feed\.json(?:\.gz)?$|feed\.\d+\.json(?:\.gz)?$")

EVENT_DATE_RE = re.compile(
    r"1\.7\.\s*Дата наступления события[^:]*:\s*(\d{1,2})\.(\d{1,2})\.(\d{4})")
SIGNATURE_RE = re.compile(r"\s3\.\s*Подпись\b.*$", re.S)
# Блок контактов в пресс-релизах («За более детальной информацией … обращайтесь
# по указанным ниже контактам: <ФИО, должности, почты>») — до ближайшего
# «Адреса страниц», пункта 2.2 или подписи.
CONTACTS_RE = re.compile(
    r"(За (?:более )?(?:детальной|подробной|дополнительной) информацией|"
    r"обращайтесь по указанным ниже контактам|Контактная информация|Контакты для СМИ|"
    r"Контакты:)"
    r".*?(?=Адреса? страниц|\s2\.2\.|\s3\.\s*Подпись|$)", re.S | re.I)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# Телефон: +7 или 8 и десять цифр с разделителями; не часть числа (не после цифры,
# буквы, точки или запятой — иначе «86043,891933…» в суммах был бы телефоном).
PHONE_RE = re.compile(r"(?<![\w+.,])(?:\+7|8)[\s\-]?\(?\d{3}\)?[\s\-]?\d{3}[\s\-]?"
                      r"\d{2}[\s\-]?\d{2}(?![\w.,]\d|\w)")


def flatten(detail: str) -> str:
    """Разметка тела сообщения → плоский текст."""
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", detail))).strip()


def clean_fact_text(detail: str) -> str:
    """Текст сообщения БЕЗ блока подписанта и контактов — то, что ложится на диск.

    Задокументированное исключение из «сырое до разбора» (D15, раздел 1
    DESIGN): ФИО, должности и почты сотрудников из сообщений в состояние не
    пишутся. Вырезаются: блок «3. Подпись …» до конца, блок контактов
    пресс-релиза; почты и телефоны — везде, последней страховкой. Разбор
    (дата события, класс) идёт по очищенному тексту — в вырезанном его нет.
    """
    text = flatten(detail)
    text = SIGNATURE_RE.sub("", " " + text).strip()
    text = CONTACTS_RE.sub("[контакты вырезаны]", text)
    text = EMAIL_RE.sub("[почта вырезана]", text)
    text = PHONE_RE.sub("[телефон вырезан]", text)
    return text.strip()


def event_date(text: str) -> str | None:
    """Дата наступления события — п. 1.7 сообщения (ДД.ММ.ГГГГ → ISO)."""
    match = EVENT_DATE_RE.search(text)
    if not match:
        return None
    day, month, year = (int(g) for g in match.groups())
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return None


def classify_disclosure(name: str) -> str:
    """Класс существенного факта по его официальному названию."""
    lowered = name.lower()
    for kind, needles in DISCLOSURE_KINDS.items():
        if any(needle.lower() in lowered for needle in needles):
            return kind
    return "other"


def refine_kind(kind: str, text: str) -> str:
    """Класс по ТЕЛУ — для «существенного влияния» и «прочего»; у решений органов
    управления — только дивиденды (о них решает совет и собрание)."""
    content = text.lower()
    marker = content.find("2. содержание сообщения")
    if marker >= 0:
        content = content[marker:]
    if kind == "board":
        return "dividends" if "дивиденд" in content else kind
    if kind not in REFINED_KINDS:
        return kind
    headline = content[len("2. содержание сообщения"):][:HEADLINE_CHARS] if marker >= 0         else content[:HEADLINE_CHARS]
    for window, order in ((headline, HEADLINE_ORDER), (content, BODY_ORDER)):
        for refined in order:
            if any(needle in window for needle in BODY_NEEDLES[refined]):
                return refined
    return kind


def event_id(link: str) -> str | None:
    match = EVENT_ID_RE.search(str(link or ""))
    return match.group(1) if match else None


def feed_items(app: dict) -> list[dict]:
    """Лента эмитента из `App.components.<feed>` → отсортированные записи.

    Запись — {id, published (день по Москве), stamp, name, link, pdf, kind}.
    Запись без id, даты или названия пропускается (её не с чем сравнивать).
    """
    items = _component(app, issuer.DISCLOSURE_FEED)
    if not isinstance(items, list):
        raise AppBlockError("лента " + ".".join(issuer.DISCLOSURE_FEED) + " — не список")
    out = []
    for item in items:
        if not isinstance(item, dict):
            continue
        identifier, day = event_id(item.get("link")), msk_day(item.get("date"))
        name = str(item.get("name") or "").strip()
        if not identifier or not day or not name:
            continue
        out.append(dict(id=identifier, published=day, stamp=int(item["date"]), name=name,
                        link=str(item.get("link")), pdf=str(item.get("pdf") or ""),
                        kind=classify_disclosure(name)))
    return sorted(out, key=lambda x: (x["stamp"], int(x["id"])))


def _raw_files(store, source: str, pattern: re.Pattern) -> list[Path]:
    directory = store.raw / source if store is not None else None
    if directory is None or not directory.is_dir():
        return []
    return sorted(item for day in directory.iterdir() if day.is_dir()
                  for item in day.iterdir() if pattern.match(item.name))


def saved_event_ids(store) -> set[str]:
    """id сообщений, чьи тела УЖЕ лежат в архиве источника (по всему архиву)."""
    return {EVENT_FILE.match(p.name).group(1)
            for p in _raw_files(store, DISCLOSURE_SOURCE, EVENT_FILE)}


def previous_feed(store) -> tuple[list[dict] | None, str]:
    """Последний сохранённый снимок ленты и его sha256 (нет — (None, ""))."""
    files = _raw_files(store, DISCLOSURE_SOURCE, FEED_FILE)
    if not files:
        return None, ""
    # Последний по дню каталога и по номеру копии внутри дня.
    def order(path: Path):
        number = re.search(r"feed\.(\d+)\.json", path.name)
        return (path.parent.name, int(number.group(1)) if number else 0)
    latest = max(files, key=order)
    body = Store.read_raw(latest)
    try:
        return json.loads(body), hashlib.sha256(body).hexdigest()
    except ValueError:
        return None, ""


def collect_lenta_disclosure(*, store=None, today: date | None = None) -> Collected:
    """Лента существенных фактов МКПАО «Лента» (lentagroup.ru, шаблон Zebra).

    Порт `collect_magnit_disclosure` 850oa с тремя отличиями Ленты:

    * **новизна — по множеству id**, а не по дате: id сообщения из ссылки
      сравнивается с последним сохранённым снимком ленты; ряд
      `<префикс>.disclosure.new` — сколько id появилось с прошлого снимка
      (первый снимок все события принимает как известные — «новых» 0);
    * **дата события — по п. 1.7 текста** сообщения («Дата наступления
      события»), а не по метке публикации: у заседания совета 16.09.2026
      публикация 17.09. Ряд `<префикс>.disclosure.events` — по дате события,
      из прочитанных тел; ряды классов `<префикс>.disclosure.<класс>` — по дню
      ПУБЛИКАЦИИ из ленты целиком (полные с первого дня);
    * **без блока подписанта и контактов** — тело ложится на диск очищенным
      (`clean_fact_text`, sha256 от очищенных байтов).

    В архив кладётся сам список ленты (≈75 КБ), а не страница (≈1 МБ
    вёрстки), и только когда он изменился. Непонятая страница —
    `AppBlockError`; страница при этом ложится на диск до отказа. Лента
    восполнима (весь архив — на одной странице), поэтому источник не в
    `IRRECOVERABLE`, но в `PROTECTED_SOURCES` хранилища (тела не ротируются).
    """
    today_iso = (today or datetime.now(timezone.utc).date()).isoformat()
    sink = RawSink(store, DISCLOSURE_SOURCE) if store is not None else None
    raw: list = []
    response = fetch(issuer.DISCLOSURE_URL)
    try:
        items = feed_items(app_json(response.text))
    except AppBlockError:
        if sink is not None:
            sink.keep("regulatory_filings.html", response)
        else:
            raw.append(("regulatory_filings.html", response.body, response))
        raise

    before, before_sha = previous_feed(store)
    feed_bytes = json.dumps(items, ensure_ascii=False, sort_keys=True).encode("utf-8")
    feed_sha = hashlib.sha256(feed_bytes).hexdigest()
    if feed_sha != before_sha:
        if sink is not None:
            sink.keep("feed.json", response, body=feed_bytes, sha256=feed_sha,
                      extra=dict(feed=".".join(issuer.DISCLOSURE_FEED), items=len(items)))
        else:
            raw.append(("feed.json", feed_bytes, response, dict(sha256=feed_sha)))

    series: dict[str, list[Point]] = {}
    stamp = dict(fetched_at=response.fetched_at, source_sha256=feed_sha)
    if items:
        last = items[-1]
        series[issuer.series("disclosure.latest")] = [Point(
            period=last["published"], value=1.0, note=last["name"][:160], **stamp)]
    by_kind_day: dict[tuple[str, str], list[str]] = {}
    for item in items:
        by_kind_day.setdefault((item["kind"], item["published"]), []).append(item["name"])
    for (kind, day), names in sorted(by_kind_day.items()):
        series.setdefault(issuer.series("disclosure." + kind), []).append(Point(
            period=day, value=float(len(names)), note="; ".join(names)[:160], **stamp))

    known = None if before is None else {str(x.get("id")) for x in before}
    new = [] if known is None else [x for x in items if x["id"] not in known]
    series[issuer.series("disclosure.new")] = [Point(
        period=today_iso, value=float(len(new)),
        note=("первый снимок ленты: %d событий приняты как известные" % len(items)
              if known is None else "; ".join(x["name"] for x in new)[:160] or "новых нет"),
        **stamp)]

    # --- тела сообщений: дата события (п. 1.7) и класс по тексту
    saved = saved_event_ids(store)
    horizon = (date.fromisoformat(today_iso) - timedelta(days=BODY_BACKFILL_DAYS)).isoformat()
    queue = [x for x in reversed(items) if x["id"] not in saved and x["published"] >= horizon]
    new_ids = {x["id"] for x in new}
    queue.sort(key=lambda x: (x["id"] not in new_ids, -x["stamp"]))
    events: list[dict] = []
    for item in queue[:MAX_EVENT_BODIES]:
        try:
            page = fetch(issuer.IR_SITE + item["link"])
        except FetchError:
            continue
        try:
            detail = str(_component(app_json(page.text), ("regulatory-filings-detail",))
                         .get("detail") or "")
        except AppBlockError:
            detail = ""
        text = clean_fact_text(detail) if detail else ""
        fact = dict(id=item["id"], link=item["link"], published=item["published"],
                    title=item["name"][:300], event_date=event_date(text),
                    kind=refine_kind(item["kind"], text), text=text,
                    parsed=bool(text))
        body = json.dumps(fact, ensure_ascii=False, sort_keys=True).encode("utf-8")
        digest = hashlib.sha256(body).hexdigest()
        extra = dict(cleaned=True, event_date=fact["event_date"], event_kind=fact["kind"])
        name = f"event_{item['id']}.json"
        if sink is not None:
            sink.keep(name, page, body=body, sha256=digest, extra=extra)
        else:
            raw.append((name, body, page, dict(sha256=digest, extra=extra)))
        saved.add(item["id"])
        if text:
            events.append(fact)

    by_event_day: dict[str, list[dict]] = {}
    for fact in events:
        by_event_day.setdefault(fact["event_date"] or fact["published"], []).append(fact)
    for day, facts in sorted(by_event_day.items()):
        series.setdefault(issuer.series("disclosure.events"), []).append(Point(
            period=day, value=float(len(facts)),
            note="; ".join(f"{f['kind']}: {f['title']}" for f in facts)[:200], **stamp))

    return Collected(DISCLOSURE_SOURCE, 7, series, raw,
                     note="событий %d, новых %d, прочитано тел %d"
                          % (len(items), len(new), len(events)))


# ------------------------------------------------------------- датабук


DATABOOK_SOURCE = "lenta_databook"
DATABOOK_SHEET = "Financials quarterly"
# Датабук на старой ссылке перепроверяется раз в неделю: компания может
# заменить файл без смены ссылки (пересчитанная история), а скачивать 1,4 МБ
# каждый такт незачем.
DATABOOK_RECHECK_DAYS = 7
DATABOOK_UNIT = 1e-3
"""Лист «Financials quarterly» — млн ₽; ряды — млрд ₽."""

QUARTER_LABEL_RE = re.compile(r"^\s*([1-4])Q\s*(\d{4})\s*$")

# Строки квартального P&L (блок IAS 17) → имя ряда `<префикс>.databook.q.<имя>`.
# Метка сравнивается без регистра и лишних пробелов; берётся ПЕРВОЕ вхождение
# (ниже на листе метки повторяются в разделе структуры SG&A и в отчёте о
# движении денег). Чистые проценты пишутся расходом со знаком «плюс» — как
# их читает ретро-проверка (`indicators.retro`).
DATABOOK_ROWS = {
    "total sales": "revenue",
    "retail sales": "revenue.retail",
    "hypermarkets": "revenue.hyper",
    "supermarkets": "revenue.super",
    "convenience stores": "revenue.conv",
    "utkonos": "revenue.utkonos",
    "drogerie": "revenue.droge",
    "remi": "revenue.remi",
    "dom lenta": "revenue.diy",
    "other formats": "revenue.other",
    "wholesale": "revenue.wholesale",
    "gross profit": "gross_profit",
    "payroll and related taxes": "payroll",
    "lease expenses": "lease",
    "ebitdar": "ebitdar_pre16",
    "ebitda": "ebitda_pre16",
    "net interest expense": "net_interest_pre16",
    "profit before income tax": "pbt_pre16",
    "net income": "net_income_pre16",
}
DATABOOK_SIGN = {"net_interest_pre16": -1.0}


class DatabookLayoutError(RuntimeError):
    """Лист датабука, из которого квартальный блок IAS 17 не вынимается."""


def databook_series_id(name: str) -> str:
    return issuer.series(f"databook.q.{name}")


def _label(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def parse_databook(body: bytes) -> dict[str, dict[str, float]]:
    """Квартальный P&L IAS 17 из датабука: {имя ряда: {квартал: млрд ₽}}.

    Шапка — строка, где встречается «IAS 17»; строка под ней — метки
    кварталов («1Q 2020»); блок кончается там, где начинается «IFRS 16» (или
    метка перестаёт быть кварталом). Пустая ячейка — нет числа, а не ноль.
    Маржа EBITDA считается здесь же (EBITDA / выручка).
    """
    import io

    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — любой отказ разбора = неверный формат
        raise DatabookLayoutError(f"файл не открывается как xlsx: {type(exc).__name__}") from exc
    if DATABOOK_SHEET not in workbook.sheetnames:
        raise DatabookLayoutError(f"нет листа {DATABOOK_SHEET!r} (листы: "
                                  + ", ".join(workbook.sheetnames) + ")")
    rows = list(workbook[DATABOOK_SHEET].iter_rows(values_only=True))
    workbook.close()
    header = next((i for i, row in enumerate(rows)
                   if any(_label(v) == "ias 17" for v in row)), None)
    if header is None or header + 1 >= len(rows):
        raise DatabookLayoutError("на листе нет шапки блока «IAS 17»")
    start = next(j for j, v in enumerate(rows[header]) if _label(v) == "ias 17")
    stop = next((j for j, v in enumerate(rows[header]) if j > start and _label(v) == "ifrs 16"),
                len(rows[header]))
    quarters: dict[int, str] = {}
    for column in range(start, stop):
        cell = rows[header + 1][column] if column < len(rows[header + 1]) else None
        match = QUARTER_LABEL_RE.match(str(cell or ""))
        if not match:
            if quarters:
                break
            continue
        quarters[column] = f"{match.group(2)}Q{match.group(1)}"
    if not quarters:
        raise DatabookLayoutError("под шапкой «IAS 17» нет меток кварталов («1Q 2020»)")

    out: dict[str, dict[str, float]] = {}
    for row in rows[header + 2:]:
        name = DATABOOK_ROWS.get(_label(row[0] if row else None))
        if name is None or name in out:
            continue
        sign = DATABOOK_SIGN.get(name, 1.0)
        values = {}
        for column, quarter in quarters.items():
            value = row[column] if column < len(row) else None
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                values[quarter] = round(sign * float(value) * DATABOOK_UNIT, 9)
        out[name] = values
    if "revenue" not in out or "ebitda_pre16" not in out:
        raise DatabookLayoutError("в блоке IAS 17 нет строк Total Sales и EBITDA")
    out["ebitda_margin_pre16"] = {
        q: round(out["ebitda_pre16"][q] / out["revenue"][q], 9)
        for q in out["revenue"] if q in out["ebitda_pre16"] and out["revenue"][q]}
    return out


def databook_link(app: dict) -> dict:
    """Ссылка на датабук со страницы публикаций: `publications.results.databook[0]`."""
    items = _component(app, ("publications", "results", "databook"))
    if not isinstance(items, list) or not items or not isinstance(items[0], dict):
        raise AppBlockError("publications.results.databook пуст")
    item = items[0]
    file = item.get("file") or {}
    link = str(file.get("link") or "")
    if not link.startswith("/"):
        raise AppBlockError("у датабука нет ссылки на файл")
    return dict(link=link, name=str(file.get("name") or ""), title=str(item.get("name") or ""),
                published=msk_day(item.get("activeFrom")))


def databook_versions(store) -> list[dict]:
    """Что архив знает о датабуке: сохранённые версии и отметки перепроверки.

    Запись — {kind: version|recheck, sha256 (файла датабука), link, day}. По
    версиям решается «эта уже есть», по всем записям — когда ссылку
    проверяли в последний раз.
    """
    out = []
    directory = store.raw / DATABOOK_SOURCE if store is not None else None
    if directory is None or not directory.is_dir():
        return out
    for meta_path in sorted(directory.glob("*/databook_*.meta.json")):
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        recheck = meta_path.name.startswith("databook_recheck")
        out.append(dict(kind="recheck" if recheck else "version",
                        sha256=meta.get("checked_sha256" if recheck else "sha256", ""),
                        link=meta.get("link", ""), day=meta_path.parent.name,
                        fetched_at=meta.get("fetched_at", ""),
                        name=meta.get("file_name", "")))
    return out


def collect_lenta_databook(*, store=None, today: date | None = None) -> Collected:
    """Датабук Ленты: сигнал «вышел отчёт» и квартальный P&L IAS 17 с винтажами.

    **Защищённый источник.** Датабук переписывает историю при каждом релизе,
    а старая ссылка со страницы исчезает: версия, которую не скачали, пропала.
    Поэтому каждая НОВАЯ версия (новый sha256) ложится в архив целиком
    (`databook_<sha12>.xlsx`, мета — sha256, ссылка, имя файла), источник — в
    `IRRECOVERABLE` и `PROTECTED_SOURCES`. Скачивание — только при новой
    ссылке или раз в `DATABOOK_RECHECK_DAYS` на прежней (замена файла без
    смены ссылки); известная версия второй раз не сохраняется.

    Ряды:
    * `<префикс>.databook.published` — день публикации (activeFrom, по Москве),
      значение 1, в примечании — название релиза и имя файла; новая точка =
      новый отчёт на странице;
    * `<префикс>.databook.new_version` — день получения новой версии (sha256
      в примечании) — сигнал такту «вышел отчёт»;
    * `<префикс>.databook.q.<имя>` — квартальные ряды IAS 17 (млрд ₽; маржа —
      доля), винтаж — момент получения файла: пересчитанная история ложится
      рядом с прежней (`Store.upsert`).
    """
    today = today or datetime.now(timezone.utc).date()
    # День каталога — `today` такта: по нему считается неделя до перепроверки.
    sink = (RawSink(store, DATABOOK_SOURCE, day=today.isoformat())
            if store is not None else None)
    raw: list = []
    page = fetch(issuer.PUBLICATIONS_URL)
    try:
        info = databook_link(app_json(page.text))
    except AppBlockError:
        if sink is not None:
            sink.keep("publications.html", page)
        else:
            raw.append(("publications.html", page.body, page))
        raise

    series: dict[str, list[Point]] = {}
    page_stamp = dict(fetched_at=page.fetched_at, source_sha256=page.sha256)
    if info["published"]:
        series[issuer.series("databook.published")] = [Point(
            period=info["published"], value=1.0,
            note=f"{info['title']}; {info['name']}"[:200], **page_stamp)]

    versions = databook_versions(store)
    known_sha = {v["sha256"] for v in versions if v["kind"] == "version"}
    same_link = [v for v in versions if v["link"] == info["link"]]
    if same_link:
        last_day = date.fromisoformat(max(v["day"] for v in same_link))
        if (today - last_day).days < DATABOOK_RECHECK_DAYS:
            return Collected(DATABOOK_SOURCE, 7, series, raw,
                             note="датабук на странице тот же; перепроверка файла — "
                                  f"не раньше чем через {DATABOOK_RECHECK_DAYS} дн")

    response = fetch(issuer.IR_SITE + info["link"], timeout=90.0)
    sha = response.sha256
    extra = dict(link=info["link"], file_name=info["name"], title=info["title"],
                 published=info["published"])
    if sha in known_sha:
        # Файл перепроверен и не изменился: второй копии в архиве не нужно,
        # но отметка перепроверки нужна — иначе неделя считалась бы от старой
        # даты, и файл качался бы каждый такт.
        marker = json.dumps(dict(sha256=sha, link=info["link"]), sort_keys=True).encode("utf-8")
        if store is not None:
            store.save_raw(DATABOOK_SOURCE, "databook_recheck.json", marker, url=response.url,
                           fetched_at=response.fetched_at, day=today.isoformat(),
                           sha256=hashlib.sha256(marker).hexdigest(),
                           extra=dict(link=info["link"], checked_sha256=sha))
        return Collected(DATABOOK_SOURCE, 7, series, raw,
                         note=f"датабук перепроверен: версия {sha[:12]} не изменилась")

    name = f"databook_{sha[:12]}.xlsx"
    if sink is not None:
        sink.keep(name, response, sha256=sha, extra=extra)
    else:
        raw.append((name, response.body, response, dict(sha256=sha, extra=extra)))
    blocks = parse_databook(response.body)
    stamp = dict(fetched_at=response.fetched_at, source_sha256=sha)
    last_quarter = max(blocks["revenue"], key=lambda q: (q[:4], q[5:])) if blocks["revenue"] else ""
    for key, values in blocks.items():
        series[databook_series_id(key)] = [
            Point(period=q, value=v, note=f"датабук {info['name']}", **stamp)
            for q, v in sorted(values.items())]
    series[issuer.series("databook.new_version")] = [Point(
        period=today.isoformat(), value=1.0,
        note=f"{info['name']}, sha256 {sha[:12]}, последний квартал {last_quarter}"[:200],
        **stamp)]
    return Collected(DATABOOK_SOURCE, 7, series, raw,
                     note=f"новая версия датабука {sha[:12]} (последний квартал {last_quarter})")


# ------------------------------------------- вакансии «Работы России»


TRUDVSEM_SOURCE = "trudvsem_vacancies"
TRUDVSEM_URL = ("http://opendata.trudvsem.ru/api/v1/vacancies/company/inn/{inn}"
                "?offset={page}&limit={limit}")
# `offset` — НОМЕР СТРАНИЦЫ, а не записи (research/06 §2.3); `limit` ≤ 100.
TRUDVSEM_LIMIT = 100
TRUDVSEM_PAGES_PER_INN = 6
TRUDVSEM_PAGES_PER_RUN = 12
TRUDVSEM_TIMEOUT = 60.0      # сервер отвечает 10–22 с
VACANCY_FILE = re.compile(r"vacancies__inn_(\d{10,12})_p(\d+)\.json(?:\.gz)?$")

# Поля вакансии, которые ЛОЖАТСЯ НА ДИСК. Список разрешённых, а не
# запрещённых: `contact_person` (ФИО), `contact_list` (телефон, почта) и
# свободные тексты (`duty`, `requirement`, `addresses`), где встречаются
# телефоны и имена, отбрасываются до записи. Задокументированное исключение из
# «сырое до разбора» (D15): sha256 в мете — от ОЧИЩЕННЫХ байтов.
VACANCY_KEEP = ("id", "creation-date", "date_modify", "salary_min", "salary_max", "currency",
                "job-name", "typicalPosition", "code_profession", "schedule", "work_places",
                "region", "company", "category", "vac_url")
VACANCY_KEEP_NESTED = {
    "region": ("region_code", "name"),
    "company": ("inn", "ogrn", "kpp", "name"),
    "category": ("specialisation",),
}


def _scrub(value):
    """Страховка поверх белого списка: почты и телефоны в строках — вырезаются."""
    if isinstance(value, str):
        return PHONE_RE.sub("[телефон вырезан]", EMAIL_RE.sub("[почта вырезана]", value))
    return value


def clean_vacancy(vacancy: dict) -> dict:
    """Вакансия без персональных данных: только поля `VACANCY_KEEP`."""
    out = {}
    for key in VACANCY_KEEP:
        if key not in vacancy:
            continue
        value = vacancy[key]
        if key in VACANCY_KEEP_NESTED:
            value = ({k: _scrub(value[k]) for k in VACANCY_KEEP_NESTED[key] if k in value}
                     if isinstance(value, dict) else None)
        else:
            value = _scrub(value)
        out[key] = value
    return out


def clean_vacancies_response(payload: dict) -> dict:
    """Ответ API → то, что ложится на диск: статус, meta и очищенные вакансии."""
    results = payload.get("results") or {}
    vacancies = results.get("vacancies") if isinstance(results, dict) else None
    cleaned = [{"vacancy": clean_vacancy(item.get("vacancy") or {})}
               for item in (vacancies or []) if isinstance(item, dict)]
    return {"status": payload.get("status"), "meta": payload.get("meta") or {},
            "results": {"vacancies": cleaned}}


def collect_trudvsem_vacancies(*, store=None) -> Collected:
    """Вакансии юрлиц группы на «Работе России» — справочный ряд вилок (D15).

    По каждому ИНН (`issuer.employer_inns()`, сверка ИНН — в `issuer`)
    страницы по 100 вакансий, пока не кончится `meta.total`, не больше
    `TRUDVSEM_PAGES_PER_INN` на ИНН и `TRUDVSEM_PAGES_PER_RUN` за такт. Каждая
    страница ОЧИЩАЕТСЯ до записи (`clean_vacancies_response`) — ФИО и контакты
    на диск не попадают.

    Ряды (день — UTC): `<префикс>.vacancies.total.<ИНН>` (`meta.total`),
    `<префикс>.vacancies.total` (сумма по опрошенным ИНН),
    `<префикс>.vacancies.with_salary` (вакансий с вилкой на прочитанных
    страницах). Цепной индекс вилок считает `indicators.salary_index` после
    сбора. Класс — `MAY_PARSE_NOTHING` и `ONCE_A_DAY` (D15: до месяца
    устойчивой работы с VPS тревоги нет); архив защищён от ротации.

    Отказ части ИНН — `Collected.error`: собранное сохраняется, статус
    источника — «ОШИБКА». Отказ всех — исключение.
    """
    sink = RawSink(store, TRUDVSEM_SOURCE) if store is not None else None
    raw: list = []
    today = UTC_TODAY()
    series: dict[str, list[Point]] = {}
    totals: dict[str, int] = {}
    with_salary = pages = 0
    failures: list[str] = []
    last: Response | None = None
    for inn in issuer.employer_inns():
        page = 0
        while page < TRUDVSEM_PAGES_PER_INN and pages < TRUDVSEM_PAGES_PER_RUN:
            url = TRUDVSEM_URL.format(inn=inn, page=page, limit=TRUDVSEM_LIMIT)
            try:
                response = fetch(url, timeout=TRUDVSEM_TIMEOUT, attempts=2)
                payload = response.json()
            except (FetchError, ValueError) as exc:
                failures.append(f"{inn} стр. {page}: {type(exc).__name__}")
                break
            pages += 1
            cleaned = clean_vacancies_response(payload)
            body = json.dumps(cleaned, ensure_ascii=False, sort_keys=True).encode("utf-8")
            name = f"vacancies__inn_{inn}_p{page}.json"
            extra = dict(cleaned=True, inn=inn, page=page,
                         removed="все поля вне белого списка (контакты, тексты, адреса)")
            if sink is not None:
                sink.keep(name, response, body=body, sha256=hashlib.sha256(body).hexdigest(),
                          extra=extra)
            else:
                raw.append((name, body, response,
                            dict(sha256=hashlib.sha256(body).hexdigest(), extra=extra)))
            last = response
            vacancies = cleaned["results"]["vacancies"]
            total = int((cleaned["meta"] or {}).get("total") or 0)
            totals[inn] = total
            with_salary += sum(1 for v in vacancies
                               if (v["vacancy"].get("salary_min") or 0) > 0
                               or (v["vacancy"].get("salary_max") or 0) > 0)
            page += 1
            if page * TRUDVSEM_LIMIT >= total or not vacancies:
                break
    if not totals:
        raise FetchError(TRUDVSEM_URL.split("?")[0], None,
                         "ни один ИНН не опрошен: " + "; ".join(failures))
    stamp = dict(fetched_at=last.fetched_at, source_sha256=last.sha256)
    for inn, total in totals.items():
        series[issuer.series(f"vacancies.total.{inn}")] = [Point(
            period=today, value=float(total), **stamp)]
    series[issuer.series("vacancies.total")] = [Point(
        period=today, value=float(sum(totals.values())),
        note=f"ИНН опрошено {len(totals)} из {len(issuer.employer_inns())}, страниц {pages}",
        **stamp)]
    series[issuer.series("vacancies.with_salary")] = [Point(
        period=today, value=float(with_salary), **stamp)]
    return Collected(TRUDVSEM_SOURCE, 2, series, raw,
                     note=f"страниц {pages}, вакансий всего {sum(totals.values())}",
                     error=("не опрошены: " + "; ".join(failures)) if failures else "")


# ------------------------------------------------------------- реестр

COLLECTORS: dict[str, tuple[Callable[[], Collected], str, int]] = {
    "cbr_key_rate": (collect_cbr_key_rate, "ежедневно", 1),
    "moex_curve": (collect_moex_curve, "ежедневно", 1),
    "moex_quote": (collect_moex_quote, "ежедневно", 1),
    "moex_security": (collect_moex_security, "ежедневно", 1),
    "bonds": (collect_bonds, "ежедневно", 1),
    "lenta_disclosure": (collect_lenta_disclosure, "ежедневно", 7),
    "lenta_databook": (collect_lenta_databook, "ежедневно", 7),
    "trudvsem_vacancies": (collect_trudvsem_vacancies, "ежедневно", 2),
}

# Каждый сборщик реестра стоит в расписании: источник, который больше не
# нужен, удаляется из реестра, а не остаётся в нём без расписания. Оба такта
# (утренний `collect` и вечерний `daily`, `ops/run.sh`) зовут `daily`: второй
# проход дня ловит отчёт, вышедший в 10:00 МСК (раскрытие и датабук), а
# сборщики раз в сутки (`ONCE_A_DAY`) второй раз источник не трогают.
DAILY = ("cbr_key_rate", "moex_curve", "moex_quote", "moex_security", "bonds",
         "lenta_disclosure", "lenta_databook", "trudvsem_vacancies")
WEEKLY: tuple[str, ...] = ()

# Без чего выпуск не выйдет, и всё остальное (аудит 850oa, вторая итерация,
# D1: отказ второстепенного источника не останавливает публикацию оценки).
# Критическими названы ровно те ряды, без которых выпуск теряет смысл:
# котировка эмитента (сравнение с рынком) и бескупонная кривая (диагностика
# устаревания книги).
CRITICAL = frozenset({"moex_quote", "moex_curve"})

# Сборщики, которым нечего отдать — и это НОРМА, а не сбой. Всем прочим
# «разобрано 0 рядов» означает мусор в ответе. «Работа России» — до месяца
# устойчивой работы с VPS (D15): с финского IP API ещё не проверялся
# (research/06 §7), и пустой ответ там — вопрос доступа, а не тревога.
MAY_PARSE_NOTHING: frozenset[str] = frozenset({"trudvsem_vacancies"})

# Сборщики, которые ходят по источнику ОДИН РАЗ В СУТКИ (признак — маркер
# `mark_collected_today` после статуса «ок»). Причина у каждого своя —
# `ONCE_A_DAY_REASONS`; «уже лежит на диске» признаком не служит (у 850oa
# обход, упавший на середине, иначе закрывал дату навсегда).
ONCE_A_DAY_REASONS = {
    "moex_security": "карточка бумаги меняется раз в годы — второй запрос за день ничего "
                     "не приносит",
    "trudvsem_vacancies": "обход ≈10 медленных страниц (10–22 с каждая) — бюджет «Работы "
                          "России» ≤ 12 страниц в сутки",
}
ONCE_A_DAY: frozenset[str] = frozenset(ONCE_A_DAY_REASONS)

# Источники, пропуск которых НЕВОСПОЛНИМ, — их отказ обязан дойти до владельца
# кодом возврата юнита (`ExecStopPost` → мост тревог). Датабук: новая версия
# заменяет старую на странице, и версия, не скачанная до замены, пропала.
# «Работа России» (история у источника не хранится) войдёт сюда решением
# владельца после месяца устойчивой работы (D15); лента раскрытия восполнима —
# весь архив на одной странице.
IRRECOVERABLE: frozenset[str] = frozenset({"lenta_databook"})

DONE_MARKER = "done_%s.json"


def done_marker_path(store, name: str, day: str) -> Path:
    return store.raw / name / day / (DONE_MARKER % name)


def already_collected_today(store, name: str, day: str | None = None) -> bool:
    """Собран ли сборщик раз-в-сутки за эту дату (есть маркер «ок»)."""
    path = done_marker_path(store, name, day or UTC_TODAY())
    if not path.exists():
        return False
    try:
        return bool(json.loads(path.read_text(encoding="utf-8")).get("complete"))
    except (OSError, ValueError):
        return False          # неразборчивый маркер равен отсутствующему


def mark_collected_today(store, name: str, day: str | None = None, *, status: str = "") -> Path:
    path = done_marker_path(store, name, day or UTC_TODAY())
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"complete": True, "status": status}, ensure_ascii=False),
                    encoding="utf-8")
    return path


def run_collectors(names, store: Store) -> dict[str, str]:
    """Запускает сборщики, складывает сырое и нормализованное.

    Падение одного источника не останавливает остальные: отчёт возвращает
    статус по каждому, а тревога поднимается наверх. Идемпотентно — повтор
    в тот же день не портит архив.
    """
    report: dict[str, str] = {}
    for name in names:
        collector, _, channel = COLLECTORS[name]
        if name in ONCE_A_DAY and already_collected_today(store, name):
            report[name] = "пропущен: за эту дату уже собран"
            continue
        try:
            # Сборщику, который умеет смотреть в архив, архив и передаётся:
            # так ряд ключевой сам дотягивается до нужной глубины на свежей
            # машине, без отдельной команды, которую надо вспомнить.
            if "store" in inspect.signature(collector).parameters:
                result = collector(store=store)
            else:
                result = collector()
        except Exception as exc:  # noqa: BLE001 — статус источника, не падение прогона
            report[name] = f"ОШИБКА: {type(exc).__name__}: {exc}"
            continue
        for entry in result.raw:
            # Четвёртый элемент — необязательные уточнения сборщика: своя
            # контрольная сумма (когда в файл ложится не весь ответ), свой день
            # каталога, добавки в мету. Без него поведение прежнее.
            filename, body, response = entry[:3]
            options = entry[3] if len(entry) > 3 else {}
            store.save_raw(result.source, filename, body, url=response.url,
                           fetched_at=response.fetched_at,
                           sha256=options.get("sha256", response.sha256),
                           day=options.get("day"), extra=options.get("extra"))
        total = 0
        rejected_before = len(store.rejected)
        for series_id, points in result.series.items():
            store.upsert(series_id, points, channel=result.channel,
                         label=series_id, cadence=COLLECTORS[name][1])
            total += len(points)
        rejected = store.rejected[rejected_before:]
        if result.error:
            # Отказ, случившийся ПОСЛЕ части сбора: сохранённое сохранено, а
            # статус источника — «ОШИБКА». Исключением этого не выразить: оно
            # ловится выше, до `save_raw` и `upsert`, и вместе с отказом одной
            # части терялся бы весь разобранный день.
            report[name] = f"ОШИБКА: {result.error}"
            continue
        if rejected:
            # Источник отдал NaN или бесконечность: конечные точки сохранены,
            # а отказ назван — как у любого мусора в ответе.
            report[name] = (f"ОШИБКА: неконечные значения отброшены ({len(rejected)}): "
                            + "; ".join(rejected[:3]))
            continue
        if not result.series and result.idle:
            # Пустой день, НАЗВАННЫЙ самим сборщиком (биржа ещё не торгует):
            # статус «ок» с причиной, точек нет. Отметку «собран за сутки» такой
            # ответ не ставит — следующий такт дня спросит источник снова.
            report[name] = f"ок: {result.idle}"
            continue
        if not result.series and name not in MAY_PARSE_NOTHING:
            # Ответ пришёл, а разобрать из него не удалось ничего. Это не
            # «пустой день», это сломанный разбор или изменившийся формат, и
            # «ок: рядов 0» был самым вредным видом зелёного.
            report[name] = "ОШИБКА: ответ получен, но разобрано 0 рядов"
            continue
        if not total and name not in MAY_PARSE_NOTHING:
            # То же самое, но на шаг тоньше: ряд ЕСТЬ, а точек в нём нет (у
            # 850oa так проходил мусор от ЦБ — «ок: рядов 1, точек 0»). Пустой
            # ряд — это ответ, из которого не вышло ни одного наблюдения, и
            # отличить его от спокойного дня может только сам источник: кто
            # вправе молчать, назван в `MAY_PARSE_NOTHING`.
            report[name] = (f"ОШИБКА: ответ получен, рядов {len(result.series)}, "
                            "но ни одного наблюдения")
            continue
        report[name] = f"ок: рядов {len(result.series)}, точек {total}"
        if name in ONCE_A_DAY:
            mark_collected_today(store, name, status=report[name])
    return report
